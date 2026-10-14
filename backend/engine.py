"""Pure import, export and validation functions for the Gaia IO Generator."""
from __future__ import annotations

import copy
import csv
import io
import json
import re
import uuid
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "gaia-io-generator-docs"
REFERENCE_XLSM = DOCS / "01_io_list_reference" / "GAIA-ELEC-1451-V1 - IO List.xlsm"
REFERENCE_CSV = DOCS / "02_fuxa_reference" / "SCADA_V1-devices_1_export_original.csv"
EXPECTED_CSV = DOCS / "02_fuxa_reference" / "SCADA_V1-devices_complet_attendu.csv"

TABLES = ["coils", "discrete_inputs", "input_registers", "holding_registers"]
VARIABLE_CATEGORIES = {"SCADA", "IO", "COM"}
TYPE_WORDS = {"BOOL": 1, "WORD": 1, "UINT": 1, "INT": 1, "BYTE": 1,
              "USINT": 1, "SINT": 1, "REAL": 2, "DWORD": 2, "UDINT": 2, "DINT": 2}
SCADA_SHEETS = {
    "SCADA-Coils": "coils", "SCADA-Discrete Inputs": "discrete_inputs",
    "SCADA-Input Registers": "input_registers", "SCADA-Holding Registers": "holding_registers",
}
SCADA_HEADERS = ["PLC Variable name", "Scada variable name", "PLC Table", "Modbus Address", "FUXA Address", "HMI Address", "Com", "FUXA / HMI", "Type", "Échelle", "Unité"]
EQUIPMENT_HEADERS = ["Equipment Identification", "PLC Variable name", "Scada variable name", "PLC Table", "Modbus Address", "Com", "Type", "Échelle", "Unité"]
TABLE_NAMES = {"coils": "SCADA_COILS", "discrete_inputs": "SCADA_DISCRETE_INPUTS", "input_registers": "SCADA_INPUT_REGISTERS", "holding_registers": "SCADA_HOLDING_REGISTERS"}
MEMADDRESS = {"coils": "000000", "discrete_inputs": "100000", "input_registers": "300000", "holding_registers": "400000"}
FUXA_TYPES = {"BOOL": "Bool", "WORD": "UInt16", "UINT": "UInt16", "INT": "Int16", "REAL": "Float32MLE", "DWORD": "UInt32MLE", "UDINT": "UInt32MLE", "DINT": "Int32MLE"}


def _id() -> str:
    return str(uuid.uuid4())


def _fresh_tag_id() -> str:
    return "t_" + uuid.uuid4().hex[:8] + "-" + uuid.uuid4().hex[:8]


def _tag_id(point: dict) -> str:
    if point.get("fuxa_id"):
        return point["fuxa_id"]
    raw = re.sub(r"[^0-9a-f]", "", str(point.get("id", "")).lower())
    raw = (raw + "0000000000000000")[:16]
    return f"t_{raw[:8]}-{raw[8:16]}"


def empty_state() -> dict:
    return {"settings": {"project_name": "Gaia Refinery", "document_reference": "GAIA-ELEC-1451", "version": "1", "device_id": "d_abeaaea7-b8914e5b", "device_name": "PLC", "ip": "192.168.0.10", "port": 502, "slave_id": 10, "polling": 500, "type_map": copy.deepcopy(FUXA_TYPES)}, "equipment_types": [], "equipment": [], "points": [], "fuxa_template": "", "fuxa_existing": [], "generation_logs": []}


def _read_workbook(source: bytes | None = None):
    return openpyxl.load_workbook(io.BytesIO(source), data_only=True) if source else openpyxl.load_workbook(REFERENCE_XLSM, data_only=True, keep_vba=True)


def _addr_number(value: Any) -> int | None:
    if value is None or value == "":
        return None
    try:
        return int(str(value), 0) if str(value).lower().startswith("0x") else int(float(value))
    except (ValueError, TypeError):
        return None


def _type_from_sheet(sheet: str, table: str) -> dict:
    code = sheet.split("-", 1)[0]
    protocol = "Wired I/O" if code == "DOSENSOR" else "Modbus RTU"
    word_order = "ABCD" if code == "ADM100W" else "CDAB"
    address_format = "hex" if code == "DIP1500" else "decimal"
    return {"id": _id(), "code": code, "manufacturer": {"ADM100W": "Trumeter", "DIP1500": "DIP", "DOSENSOR": ""}.get(code, ""), "model": code, "description": "Reference equipment type", "protocol": protocol, "word_order": word_order, "address_format": address_format, "defaults": {"baud": 9600, "parity": "None", "stopbits": 2 if code == "ADM100W" else 1, "slave_id": 1}, "registers": [], "point_templates": [], "modes": ([{"name": "time", "register_types": {"0x4005": "UDINT", "0x4007": "UDINT", "0x4009": "UDINT"}}, {"name": "volume", "register_types": {"0x4005": "REAL", "0x4007": "REAL", "0x4009": "REAL"}}] if code == "DIP1500" else [])}


def _excel_points(wb) -> tuple[list[dict], list[dict], list[dict]]:
    types_by_code: dict[str, dict] = {}
    eq_by_tag: dict[str, dict] = {}
    points: list[dict] = []
    for sheet_name in wb.sheetnames:
        if sheet_name.startswith("SCADA-") or sheet_name == "Navigation":
            continue
        ws = wb[sheet_name]
        if not ws.cell(1, 1).value or str(ws.cell(1, 1).value) != "Equipment Identification":
            continue
        code = sheet_name.split("-", 1)[0]
        table = next((x for x in ("coils", "discrete_inputs", "input_registers", "holding_registers") if x.replace("_", " ") in sheet_name.lower()), "input_registers")
        et = types_by_code.setdefault(code, _type_from_sheet(sheet_name, table))
        for row in ws.iter_rows(min_row=2, values_only=True):
            label = row[0]
            if not label:
                continue
            plc_table, address = row[3] if len(row) > 3 else "", row[4] if len(row) > 4 else None
            et["registers"].append({"table": table, "address": address, "label": label, "type": row[6] if len(row) > 6 else "", "words": 1, "access": row[5] if len(row) > 5 else "", "scale": row[7] if len(row) > 7 else "", "unit": row[8] if len(row) > 8 else "", "plc_table": plc_table})
    for sheet_name, table in SCADA_SHEETS.items():
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        idx = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            variable = row[0]
            if not variable:
                idx += 1
                continue
            address = _addr_number(row[3])
            if address is None:
                idx += 1
                continue
            if table == "discrete_inputs":
                address -= 10000
            plcvar = str(variable)
            match = re.search(r"(?:SCADA|IO|COM)_([A-Za-z0-9_]+_\d+)\.", plcvar)
            tag = match.group(1) if match else "DOSENSOR_1"
            category = "IO" if plcvar.startswith("GVL_IO.") else "COM" if plcvar.startswith("GVL_COM.") else "SCADA"
            code = re.sub(r"_\d+$", "", tag)
            if code not in types_by_code:
                types_by_code[code] = _type_from_sheet(code + "-Input Registers", table)
            if tag not in eq_by_tag:
                eq_by_tag[tag] = {"id": _id(), "tag": tag, "type_id": types_by_code[code]["id"], "description": "", "active": True, "status": "Existing", "communication": {}, "notes": ""}
            scada_name = row[1] or ""
            data_type = str(row[8] or "")
            pair = table in ("input_registers", "holding_registers") and data_type.upper() not in ("WORD", "UINT", "INT", "UINT16", "INT16") and idx + 1 < ws.max_row and ws.cell(idx + 3, 1).value == variable
            base_name = str(scada_name).removesuffix(" (2nd register)")
            existing = next((p for p in points if p["plc_variable"] == plcvar and p["table"] == table and p["equipment_id"] == eq_by_tag[tag]["id"]), None)
            if existing:
                idx += 1
                continue
            target = row[7] or "FUXA"
            points.append({"id": _id(), "equipment_id": eq_by_tag[tag]["id"], "plc_variable": plcvar, "scada_name": base_name, "table": table, "address": address, "locked": False, "codesys_type": data_type, "fuxa_type": FUXA_TYPES.get(data_type.upper(), ""), "words": 2 if pair else 1, "divisor": 1, "scale": row[9] or "1", "unit": row[10] or "", "direction": row[6] or "Plc To SCADA", "target": target, "hmi_address": row[5] or "", "daq_enabled": False, "daq_interval": 60, "fuxa_id": _fresh_tag_id(), "comment": ""})
            points[-1]["variable_category"] = category
            types_by_code[code]["point_templates"].append({k: points[-1].get(k) for k in ("plc_variable", "scada_name", "table", "codesys_type", "fuxa_type", "words", "divisor", "scale", "unit", "direction", "target", "hmi_address", "daq_enabled", "daq_interval", "comment")})
            types_by_code[code]["point_templates"][-1]["variable_category"] = category
            idx += 1
    return list(types_by_code.values()), list(eq_by_tag.values()), points


def seed_state() -> dict:
    state = empty_state()
    types, equipment, points = _excel_points(_read_workbook())
    state["equipment_types"], state["equipment"], state["points"] = types, equipment, points
    by_id = {t["id"]: t for t in types}
    for p in points:
        eq = next((e for e in equipment if e["id"] == p["equipment_id"]), None)
        if eq:
            typ = by_id[eq["type_id"]]
            if not any(x.get("plc_variable") == p["plc_variable"] for x in typ["point_templates"]):
                typ["point_templates"].append({k: p.get(k) for k in ("plc_variable", "scada_name", "table", "codesys_type", "fuxa_type", "words", "divisor", "scale", "unit", "direction", "target", "hmi_address", "daq_enabled", "daq_interval", "comment")})
    state["fuxa_template"] = REFERENCE_CSV.read_text(encoding="utf-8")
    return import_fuxa(REFERENCE_CSV.read_bytes(), state)


def import_excel(data: bytes, state: dict) -> dict:
    result = copy.deepcopy(state)
    wb = _read_workbook(data)
    if "_GaiaMetadata" in wb.sheetnames and wb["_GaiaMetadata"]["A1"].value:
        try:
            meta_ws = wb["_GaiaMetadata"]
            stored = json.loads("".join(str(meta_ws.cell(r, 1).value or "") for r in range(1, meta_ws.max_row + 1)))
            # Metadata preserves the full editable model, including fields not represented in Excel.
            has_current_model = any(state.get(key) for key in ("equipment_types", "equipment", "points"))
            result = copy.deepcopy(state) if has_current_model else stored
            for key in ("settings", "equipment_types", "equipment", "points", "fuxa_template", "fuxa_existing", "generation_logs"):
                if key not in result or result[key] in (None, [], {}, ""):
                    result[key] = copy.deepcopy(stored.get(key, result.get(key)))
            seen: set[tuple] = set()
            for sheet_name, table in SCADA_SHEETS.items():
                if sheet_name not in wb.sheetnames:
                    continue
                ws = wb[sheet_name]
                for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                    if not row[0] or not row[1] or _addr_number(row[3]) is None or str(row[1]).endswith(" (2nd register)"):
                        continue
                    address = _addr_number(row[3])
                    if table == "discrete_inputs": address -= 10000
                    p = next((p for p in result.get("points", []) if p.get("table") == table and p.get("plc_variable") == row[0]), None)
                    if p is None:
                        p = next((p for p in result.get("points", []) if p.get("table") == table and int(p.get("address") or -1) == address), None)
                    if p:
                        p.update({"plc_variable": row[0], "address": address, "words": 2 if row_num < ws.max_row and ws.cell(row_num + 1, 1).value == row[0] and str(ws.cell(row_num + 1, 2).value or "").endswith(" (2nd register)") else 1, "scada_name": row[1] or "", "hmi_address": row[5] or "", "direction": row[6] or "", "target": row[7] or "FUXA", "codesys_type": row[8] or "", "scale": row[9] or "", "unit": row[10] or ""})
                        if not p.get("fuxa_type_override"):
                            p["fuxa_type"] = result.get("settings", {}).get("type_map", FUXA_TYPES).get(str(row[8] or "").upper(), p.get("fuxa_type", ""))
                        seen.add((table, address, p.get("plc_variable")))
                    else:
                        match = re.search(r"SCADA_([A-Z0-9]+_\d+)\.", str(row[0]))
                        tag = match.group(1) if match else ""
                        equipment = next((e for e in result.get("equipment", []) if e.get("tag") == tag), None)
                        words = 2 if row_num < ws.max_row and ws.cell(row_num + 1, 1).value == row[0] and str(ws.cell(row_num + 1, 2).value or "").endswith(" (2nd register)") else 1
                        result.setdefault("points", []).append({"id": _id(), "equipment_id": equipment.get("id", "") if equipment else "", "plc_variable": row[0], "scada_name": row[1], "table": table, "address": address, "locked": False, "codesys_type": row[8] or "", "fuxa_type": result.get("settings", {}).get("type_map", FUXA_TYPES).get(str(row[8] or "").upper(), ""), "fuxa_type_override": False, "words": words, "divisor": 1, "scale": row[9] or "1", "unit": row[10] or "", "direction": row[6] or "", "target": row[7] or "FUXA", "hmi_address": row[5] or "", "daq_enabled": False, "daq_interval": 60, "fuxa_id": _fresh_tag_id(), "comment": ""})
                        seen.add((table, address, row[0]))
            result["points"] = [p for p in result.get("points", []) if not _is_active(result, p) or not _is_scada(p) or (p.get("table"), int(p.get("address") or -1), p.get("plc_variable")) in seen]
            # Keep catalogue-only metadata while applying edits made to visible equipment register sheets.
            for typ in result.get("equipment_types", []):
                for table in TABLES:
                    sheet_name = f"{typ.get('code')}-{table.replace('_', ' ').title()}"[:31]
                    if sheet_name not in wb.sheetnames:
                        continue
                    ws = wb[sheet_name]
                    edited = []
                    for row in ws.iter_rows(min_row=2, values_only=True):
                        if not row[0]: continue
                        address = row[4] if len(row) > 4 else None
                        old = next((r for r in typ.get("registers", []) if r.get("table") == table and str(r.get("address")) == str(address)), None)
                        if old is None:
                            old = next((r for r in typ.get("registers", []) if r.get("table") == table and str(r.get("label", "")).casefold() == str(row[0]).casefold()), {})
                        reg = copy.deepcopy(old)
                        reg.update({"table": table, "label": row[0], "plc_table": row[3] or "", "address": address, "access": row[5] or "", "type": row[6] or "", "scale": row[7] or "", "unit": row[8] or ""})
                        edited.append(reg)
                    typ["registers"] = [r for r in typ.get("registers", []) if r.get("table") != table] + edited
            return result
        except (ValueError, TypeError, KeyError):
            pass
    incoming_types, incoming_eq, incoming_points = _excel_points(wb)
    incoming_type_id_to_code = {t["id"]: t["code"] for t in incoming_types}
    incoming_eq_id_to_tag = {e["id"]: e["tag"] for e in incoming_eq}
    current_types = {x.get("code"): x for x in result.get("equipment_types", [])}
    for typ in incoming_types:
        if typ["code"] in current_types:
            typ["id"] = current_types[typ["code"]]["id"]
        else:
            current_types[typ["code"]] = typ
    type_id = {k: v["id"] for k, v in current_types.items()}
    current_eq = {x.get("tag"): x for x in result.get("equipment", [])}
    for eq in incoming_eq:
        source_type_id = eq["type_id"]
        eq["type_id"] = type_id.get(incoming_type_id_to_code.get(source_type_id, ""), source_type_id)
        prior = current_eq.get(eq["tag"])
        if prior:
            eq["id"] = prior["id"]
            eq.update({k: v for k, v in prior.items() if k not in eq})
        current_eq[eq["tag"]] = eq
    eqid_map = {e["tag"]: e["id"] for e in current_eq.values()}
    existing_points = {(p.get("plc_variable"), p.get("table")): p for p in result.get("points", [])}
    for p in incoming_points:
        tag = incoming_eq_id_to_tag.get(p["equipment_id"], "")
        p["equipment_id"] = eqid_map.get(tag, p["equipment_id"])
        p["fuxa_type_override"] = False
        p["fuxa_type"] = result.get("settings", {}).get("type_map", FUXA_TYPES).get(str(p.get("codesys_type", "")).upper(), p.get("fuxa_type", ""))
        old = existing_points.get((p["plc_variable"], p["table"]))
        if old:
            p["id"] = old["id"]
            for key in ("locked", "fuxa_id", "daq_enabled", "daq_interval", "target", "hmi_address", "comment"):
                p[key] = old.get(key, p.get(key))
    result["equipment_types"] = list(current_types.values())
    result["equipment"] = list(current_eq.values())
    # The workbook is an authoritative IO list for matching points; unmatched current points are reported by callers if desired.
    result["points"] = incoming_points + [p for p in state.get("points", []) if not _is_scada(p) and (p.get("plc_variable"), p.get("table")) not in {(x.get("plc_variable"), x.get("table")) for x in incoming_points}]
    return result


def import_fuxa(data: bytes, state: dict) -> dict:
    result = copy.deepcopy(state)
    text = data.decode("utf-8-sig")
    lines = text.splitlines(keepends=True)
    result["fuxa_template"] = text
    result["settings"] = copy.deepcopy(result.get("settings", {}))
    device_row = next((line.rstrip("\r\n").split(",") for line in lines if line.startswith("D@,") and len(line.rstrip("\r\n").split(",")) >= 18 and line.rstrip("\r\n").split(",")[4] == "ModbusTCP"), None)
    if device_row:
        result["settings"].update({"device_id": device_row[1], "device_name": device_row[2], "polling": _addr_number(device_row[5]) or 500, "ip": device_row[6], "port": _addr_number(device_row[13]) or 502, "slave_id": _addr_number(device_row[15]) or 10})
    existing = []
    for line in lines:
        if not line.startswith("T@,"):
            continue
        row = line.rstrip("\r\n").split(",")
        if len(row) < 21 or row[1] != result["settings"].get("device_id", "d_abeaaea7-b8914e5b"):
            continue
        if row[3].strip().lower().startswith("test"):
            continue
        existing.append({"device_id": row[1], "id": row[2], "name": row[3], "type": row[5], "memaddress": row[6], "address": row[7], "line": line})
    result["fuxa_existing"] = existing
    for p in result.get("points", []):
        if not _is_scada(p):
            continue
        found = next((t for t in existing if t["memaddress"] == MEMADDRESS.get(p.get("table")) and _addr_number(t["address"]) == _fuxa_address(p)), None)
        if not found:
            found = next((t for t in existing if t["name"].casefold() == str(p.get("scada_name", "")).casefold()), None)
        if found:
            p["fuxa_id"], p["scada_name"] = found["id"], found["name"]
            if found["type"]:
                p["fuxa_type"] = found["type"]
                p["fuxa_type_override"] = True
    return result


def _fuxa_address(point: dict) -> int:
    return int(point.get("address") or 0)


def _is_active(state: dict, point: dict) -> bool:
    equipment = next((e for e in state.get("equipment", []) if e.get("id") == point.get("equipment_id")), None)
    return equipment is None or equipment.get("active", True) is not False


def _is_scada(point: dict) -> bool:
    """Legacy points without an explicit category remain SCADA points."""
    return point.get("variable_category", "SCADA") == "SCADA"


def _point_fuxa_type(state: dict, point: dict) -> str:
    if point.get("fuxa_type_override"):
        return str(point.get("fuxa_type") or "")
    mapping = state.get("settings", {}).get("type_map", FUXA_TYPES)
    return str(mapping.get(str(point.get("codesys_type", "")).upper(), point.get("fuxa_type", "")))


def allocate_addresses(state: dict) -> dict:
    result = copy.deepcopy(state)
    points = result.get("points", [])
    for table in TABLES:
        occupied: set[int] = set()
        for p in points:
            if p.get("table") == table and p.get("address"):
                occupied.update(range(int(p["address"]), int(p["address"]) + int(p.get("words") or 1)))
        cursor = 1
        for p in points:
            if p.get("table") != table or p.get("address") or p.get("locked"):
                continue
            size = max(1, int(p.get("words") or 1))
            while any(n in occupied for n in range(cursor, cursor + size)):
                cursor += 1
            p["address"] = cursor
            occupied.update(range(cursor, cursor + size))
            cursor += size
    return result


def validate(state: dict) -> dict:
    errors, warnings = [], []
    if not isinstance(state, dict):
        return {"errors": [{"message": "State must be an object"}], "warnings": []}
    settings = state.get("settings")
    types, equipment, points = state.get("equipment_types"), state.get("equipment"), state.get("points")
    if not isinstance(settings, dict): errors.append({"message": "settings must be an object"}); settings = {}
    for key in ("equipment_types", "equipment", "points"):
        if not isinstance(state.get(key), list): errors.append({"message": f"{key} must be an array"})
    if not isinstance(types, list): types = []
    if not isinstance(equipment, list): equipment = []
    if not isinstance(points, list): points = []
    for key in ("project_name", "document_reference", "version", "device_id", "device_name", "ip"):
        if key in settings and not isinstance(settings[key], str): errors.append({"message": f"settings.{key} must be a string"})
    def unique_ids(items: list, label: str):
        seen_ids = set()
        for item in items:
            item_id = item.get("id") if isinstance(item, dict) else None
            if not item_id: errors.append({"message": f"{label} id is required"})
            elif item_id in seen_ids: errors.append({"message": f"Duplicate {label} id: {item_id}"})
            seen_ids.add(item_id)
    unique_ids(types, "equipment type"); unique_ids(equipment, "equipment"); unique_ids(points, "point")
    collection_ids = [item.get("id") for collection in (types, equipment, points) for item in collection if isinstance(item, dict) and item.get("id")]
    if len(collection_ids) != len(set(collection_ids)): errors.append({"message": "IDs must be unique across all collections"})
    type_by_id = {t.get("id"): t for t in types if isinstance(t, dict)}
    type_codes = set()
    for typ in types:
        if not isinstance(typ, dict): continue
        code = str(typ.get("code") or "")
        if not code or not re.fullmatch(r"[A-Za-z0-9_-]+", code): errors.append({"message": f"Invalid equipment type code: {code}"})
        if code.casefold() in type_codes: errors.append({"message": f"Duplicate equipment type code: {code}"})
        type_codes.add(code.casefold())
        if not isinstance(typ.get("registers", []), list): errors.append({"message": f"Equipment type {code} registers must be an array"})
        if not isinstance(typ.get("point_templates", []), list): errors.append({"message": f"Equipment type {code} point_templates must be an array"})
        if not isinstance(typ.get("defaults", {}), dict): errors.append({"message": f"Equipment type {code} defaults must be an object"})
        for template in typ.get("point_templates", []) if isinstance(typ.get("point_templates", []), list) else []:
            if not isinstance(template, dict):
                errors.append({"message": f"Equipment type {code} variables must be objects"})
            elif template.get("variable_category", "SCADA") not in VARIABLE_CATEGORIES:
                errors.append({"message": f"Invalid variable category in equipment type {code}"})
    equipment_by_id = {e.get("id"): e for e in equipment if isinstance(e, dict)}
    equipment_tags = set()
    for eq in equipment:
        if not isinstance(eq, dict): continue
        tag = str(eq.get("tag") or "")
        if not tag or not re.fullmatch(r"[A-Za-z0-9_-]+", tag): errors.append({"message": f"Invalid equipment tag: {tag}"})
        if tag.casefold() in equipment_tags: errors.append({"message": f"Duplicate equipment tag: {tag}"})
        equipment_tags.add(tag.casefold())
        if eq.get("type_id") not in type_by_id: errors.append({"message": f"Equipment {tag} has an unknown equipment type"})
        if "active" in eq and not isinstance(eq["active"], bool): errors.append({"message": f"Equipment {tag} active must be boolean"})
        if "communication" in eq and not isinstance(eq["communication"], dict): errors.append({"message": f"Equipment {tag} communication must be an object"})
    safe = re.compile(r'^[^,~"\r\n\x00-\x1f]+$')
    for key in ("device_id", "device_name", "ip"):
        value = str(settings.get(key, "") or "")
        unsafe_prefix = value[:1] in "=+-@"
        if not value or not safe.fullmatch(value) or value != value.strip() or unsafe_prefix: errors.append({"message": f"Invalid device {key}"})
    if settings.get("device_id") and not re.fullmatch(r"d_[A-Za-z0-9-]+", str(settings["device_id"])): errors.append({"message": "Invalid device id format"})
    for key, low, high in (("port", 1, 65535), ("slave_id", 1, 247), ("polling", 1, 2147483647)):
        try:
            value = int(settings.get(key))
            if not low <= value <= high: raise ValueError
        except (TypeError, ValueError): errors.append({"message": f"Invalid device {key}"})
    for key in ("port", "slave_id", "polling"):
        if key in settings and not isinstance(settings[key], int): errors.append({"message": f"settings.{key} must be an integer"})
    fuxa_valid = {"Bool", "Int16", "UInt16", "Int32", "UInt32", "Float32", "Float64", "Int64", "Int16LE", "UInt16LE", "Int32LE", "UInt32LE", "Float32LE", "Float64LE", "Int64LE", "Int32MLE", "UInt32MLE", "Float32MLE", "Float64MLE", "String"}
    type_map = settings.get("type_map", {})
    if not isinstance(type_map, dict): errors.append({"message": "type_map must be an object"}); type_map = {}
    for source, target_type in type_map.items():
        if not safe.fullmatch(str(source)) or target_type not in fuxa_valid: errors.append({"message": f"Invalid type mapping for {source}"})
    addresses: dict[tuple, dict] = {}
    names: dict[str, dict] = {}
    fuxa_ids: set[str] = set()
    plc_variables: set[str] = set()
    for p in points:
        if not isinstance(p, dict): errors.append({"message": "Each point must be an object"}); continue
        if p.get("variable_category", "SCADA") not in VARIABLE_CATEGORIES:
            errors.append({"message": "Variable category must be SCADA, IO, or COM", "point_id": p.get("id")})
        if p.get("equipment_id") not in equipment_by_id: errors.append({"message": f"Point {p.get('id')} has an unknown equipment_id", "point_id": p.get("id")})
        table, address = p.get("table"), p.get("address")
        if table not in TABLES:
            errors.append({"message": f"Unknown table: {table}", "point_id": p.get("id")})
            continue
        if not p.get("plc_variable"):
            errors.append({"message": "PLC variable name is required", "point_id": p.get("id")})
        variable = str(p.get("plc_variable") or "")
        if variable and (not safe.fullmatch(variable) or variable != variable.strip()): errors.append({"message": f"Invalid PLC variable name: {variable}", "point_id": p.get("id")})
        if variable.casefold() in plc_variables: errors.append({"message": f"Duplicate PLC variable name: {variable}", "point_id": p.get("id")})
        plc_variables.add(variable.casefold())
        if not p.get("scada_name"):
            errors.append({"message": "SCADA name is required", "point_id": p.get("id")})
        name = str(p.get("scada_name") or "")
        if name != name.strip() or re.search(r'[,~"\r\n\x00-\x1f]', name) or name[:1] in "=+-@":
            errors.append({"message": f"Invalid SCADA name: {name}", "point_id": p.get("id")})
        elif name:
            if name.casefold() in names:
                errors.append({"message": f"Duplicate SCADA name: {name}", "point_id": p.get("id")})
            names[name.casefold()] = p
        if p.get("target", "FUXA") not in ("FUXA", "HMI"): errors.append({"message": f"Invalid target for {name}", "point_id": p.get("id")})
        if p.get("direction") not in ("Plc To SCADA", "SCADA To Plc", "Two Ways", "Equipment To Plc", "Plc To Equipment"): errors.append({"message": f"Invalid direction for {name}", "point_id": p.get("id")})
        if p.get("fuxa_type") and p.get("fuxa_type") not in fuxa_valid: errors.append({"message": f"Invalid FUXA type for {name}", "point_id": p.get("id")})
        for flag in ("locked", "daq_enabled"):
            if flag in p and not isinstance(p[flag], bool): errors.append({"message": f"{flag} must be boolean for {name}", "point_id": p.get("id")})
        if "daq_interval" in p:
            try:
                if int(p["daq_interval"]) < 1: raise ValueError
            except (TypeError, ValueError): errors.append({"message": f"Invalid DAQ interval for {name}", "point_id": p.get("id")})
        fuxa_id = p.get("fuxa_id")
        if fuxa_id:
            if not re.fullmatch(r"t_[0-9a-f]{8}-[0-9a-f]{8}", str(fuxa_id)): errors.append({"message": f"Invalid FUXA tag id for {name}", "point_id": p.get("id")})
            if fuxa_id in fuxa_ids: errors.append({"message": f"Duplicate FUXA tag id: {fuxa_id}", "point_id": p.get("id")})
            fuxa_ids.add(fuxa_id)
        if _is_scada(p) and not _point_fuxa_type(state, p) and (p.get("target", "FUXA") == "FUXA"):
            errors.append({"message": f"FUXA type is required for {name}", "point_id": p.get("id")})
        try:
            divisor = float(p.get("divisor", 1))
            if not divisor > 0: raise ValueError
        except (TypeError, ValueError): errors.append({"message": f"Invalid divisor for {name}", "point_id": p.get("id")})
        if not p.get("unit"):
            warnings.append({"message": f"Unit is empty for {name}", "point_id": p.get("id")})
        if address is None:
            errors.append({"message": f"Address is missing for {name}", "point_id": p.get("id")})
            continue
        try:
            start, words = int(address), int(p.get("words") or 1)
        except (ValueError, TypeError):
            errors.append({"message": f"Invalid address for {name}", "point_id": p.get("id")})
            continue
        if words not in (1, 2): errors.append({"message": f"Invalid word count for {name}", "point_id": p.get("id")}); continue
        codesys_type = str(p.get("codesys_type", "")).upper()
        expected_words = TYPE_WORDS.get(codesys_type)
        if expected_words is not None and words != expected_words:
            errors.append({"message": f"{codesys_type} requires {expected_words} word(s) for {name}", "point_id": p.get("id")})
        if table in ("coils", "discrete_inputs") and codesys_type and codesys_type != "BOOL":
            errors.append({"message": f"Bit tables require BOOL for {name}", "point_id": p.get("id")})
        if str(address) != str(start) or isinstance(address, bool):
            errors.append({"message": f"Address must be a whole number for {name}", "point_id": p.get("id")})
        if start < 1 or start + words - 1 > 65536: errors.append({"message": f"Address out of range for {name}", "point_id": p.get("id")}); continue
        # HMI points still occupy the PLC's Modbus address space.
        for n in range(start, start + words):
            key = (table, n)
            if key in addresses:
                errors.append({"message": f"Address collision in {table} at {n}", "point_id": p.get("id"), "other_point_id": addresses[key].get("id")})
            else:
                addresses[key] = p
    return {"errors": errors, "warnings": warnings}


def export_fuxa(state: dict) -> bytes:
    settings = state.get("settings", {})
    template = state.get("fuxa_template") or REFERENCE_CSV.read_text(encoding="utf-8")
    lines = template.replace("\r\n", "\n").replace("\r", "\n").splitlines(keepends=True)
    original_device = next((l.rstrip("\r\n").split(",")[1] for l in lines if l.startswith("D@,") and len(l.rstrip("\r\n").split(",")) > 4 and l.rstrip("\r\n").split(",")[4] == "ModbusTCP"), settings.get("device_id", "d_abeaaea7-b8914e5b"))
    # Preserve fixed CSV sections and all original device rows. Keep imported real tags in their original order.
    existing_ids = {t.get("id") for t in state.get("fuxa_existing", [])}
    point_by_fuxa_id = {p.get("fuxa_id"): p for p in state.get("points", []) if p.get("fuxa_id")}
    retained: list[str] = []
    for line in lines:
        if not line.startswith("T@,"):
            if line.startswith("D@,"):
                row = line.rstrip("\r\n").split(",")
                if len(row) >= 18 and row[1] == original_device:
                    row[1] = settings.get("device_id", row[1]); row[2] = settings.get("device_name", row[2]); row[5] = str(settings.get("polling", row[5])); row[6] = settings.get("ip", row[6]); row[13] = "" if settings.get("port", 502) == 502 else str(settings["port"]); row[15] = str(settings.get("slave_id", row[15]))
                    line = ",".join(row) + "\n"
            retained.append(line)
            continue
        row = line.rstrip("\r\n").split(",")
        name = row[3] if len(row) > 3 else ""
        if name.strip().lower().startswith("test"):
            continue
        point = point_by_fuxa_id.get(row[2])
        if row[1] == original_device and name != "PLC Connection Status" and (not point or not _is_scada(point) or point.get("target", "FUXA") != "FUXA" or not _is_active(state, point)):
            continue
        if row[1] == original_device and row[2] not in existing_ids:
            continue
        if point:
            row = ["T@", settings.get("device_id", row[1]), row[2], point.get("scada_name", row[3]), "", _point_fuxa_type(state, point), MEMADDRESS.get(point.get("table"), row[6]), str(_fuxa_address(point)), str(point.get("divisor") or 1), "", "", "", "", "", "", "", str(bool(point.get("daq_enabled", False))).lower(), str(point.get("daq_interval", 60)), "false", "false", ""]
            line = ",".join(row) + "\n"
        elif name == "PLC Connection Status":
            row[6] = settings.get("device_id", row[6])
            line = ",".join(row) + "\n"
        retained.append(line)
    emitted_ids = {l.rstrip("\r\n").split(",")[2] for l in retained if l.startswith("T@,")}
    emitted_names = {l.rstrip("\r\n").split(",")[3].casefold() for l in retained if l.startswith("T@,")}
    additions = []
    for p in state.get("points", []):
        if not _is_scada(p) or p.get("target", "FUXA") != "FUXA" or not _is_active(state, p):
            continue
        if not _point_fuxa_type(state, p) or not p.get("scada_name"):
            continue
        fid = _tag_id(p)
        if fid in emitted_ids:
            continue
        name = str(p["scada_name"])
        if name.casefold() in emitted_names:
            continue
        emitted_ids.add(fid); emitted_names.add(name.casefold())
        words = ["T@", settings.get("device_id", "d_abeaaea7-b8914e5b"), fid, name, "", _point_fuxa_type(state, p), MEMADDRESS[p["table"]], str(_fuxa_address(p)), str(p.get("divisor") or 1), "", "", "", "", "", "", "", str(bool(p.get("daq_enabled", False))).lower(), str(p.get("daq_interval", 60)), "false", "false", ""]
        additions.append(",".join(words) + "\n")
    # Insert new tags after the final retained tag while keeping the fixed template lines and spacing.
    last_tag = max((i for i, line in enumerate(retained) if line.startswith("T@,")), default=-1)
    retained[last_tag + 1:last_tag + 1] = additions
    content = "".join(retained)
    return (content.rstrip("\n") + "\n\n\n").encode("utf-8")


def _sheets_and_rows(state: dict):
    scada: dict[str, list[list[Any]]] = {t: [] for t in TABLES}
    for p in state.get("points", []):
        table = p.get("table")
        if table not in scada or not p.get("address") or not _is_active(state, p):
            continue
        address, words = int(p["address"]), max(1, int(p.get("words") or 1))
        table_label = {"coils": "Coil", "discrete_inputs": "Discrete Input", "input_registers": "Input Register", "holding_registers": "Holding Register"}[table]
        plc_table = f"{table_label} TCP HMI {address - 1}"
        if table in ("coils", "discrete_inputs"):
            plc_table = f"{table_label} TCP HMI {divmod(address - 1, 8)[0]}.{divmod(address - 1, 8)[1]}"
        base = [p.get("plc_variable", ""), p.get("scada_name", ""), plc_table, address if table != "discrete_inputs" else 10000 + address, 100000 + address if table == "discrete_inputs" else address, p.get("hmi_address", ""), p.get("direction", "Plc To SCADA"), p.get("target", "FUXA"), p.get("codesys_type", ""), p.get("scale", "1"), p.get("unit", "")]
        scada[table].append(base)
        if words == 2:
            second = list(base); second[1] = (str(base[1]) + " (2nd register)").strip(); second[2] = f"{table_label} TCP HMI {address}"; second[3] = address + 1 if table != "discrete_inputs" else 10001 + address; second[4] = address + 1 if table != "discrete_inputs" else 100001 + address
            scada[table].append(second)
    # Preserve open address rows used by the reference PLC mapping.
    capacities = {"coils": 504, "discrete_inputs": 504, "input_registers": 501, "holding_registers": 501}
    for table, rows in scada.items():
        rows.sort(key=lambda r: r[3])
        present = {r[3] for r in rows}
        for i in range(1, capacities[table] + 1):
            address = 10000 + i if table == "discrete_inputs" else i
            if address not in present:
                table_label = {"coils": "Coil", "discrete_inputs": "Discrete Input", "input_registers": "Input Register", "holding_registers": "Holding Register"}[table]
                plc_index = i - 1
                plc_table = f"{table_label} TCP HMI {divmod(plc_index, 8)[0]}.{divmod(plc_index, 8)[1]}" if table in ("coils", "discrete_inputs") else f"{table_label} TCP HMI {plc_index}"
                fuxa_address = 100000 + i if table == "discrete_inputs" else i
                rows.append(["", "", plc_table, address, fuxa_address, "", "", "", "", "", ""])
        rows.sort(key=lambda r: r[3])
    return scada


def export_excel(state: dict) -> bytes:
    wb = openpyxl.Workbook()
    nav = wb.active; nav.title = "Navigation"
    nav.cell(1, 2, "Navigation – click a sheet")
    nav_rows = [("📁", "SCADA (general)")]
    for sheet in SCADA_SHEETS:
        nav_rows.append((None, sheet))
    for typ in state.get("equipment_types", []):
        nav_rows.append(("📁", typ.get("code", "Equipment")))
        for table in TABLES:
            if any(r.get("table") == table for r in typ.get("registers", [])):
                nav_rows.append((None, f"{typ['code']}-{table.replace('_', ' ').title()}"))
    actual_sheet_names = set(SCADA_SHEETS) | {f"{typ.get('code')}-{table.replace('_', ' ').title()}"[:31] for typ in state.get("equipment_types", []) for table in TABLES if any(r.get("table") == table for r in typ.get("registers", []))}
    for r, row in enumerate(nav_rows, start=3):
        nav.cell(r, 1, row[0]); nav.cell(r, 2, row[1])
        if row[1] in actual_sheet_names:
            nav.cell(r, 2).hyperlink = f"#'{row[1]}'!A1"
    scada = _sheets_and_rows(state)
    for sheet, table in SCADA_SHEETS.items():
        ws = wb.create_sheet(sheet)
        for c, h in enumerate(SCADA_HEADERS, 1): ws.cell(1, c, h)
        for r, row in enumerate(scada[table], 2):
            for c, v in enumerate(row, 1): ws.cell(r, c, v)
        _add_table(ws, TABLE_NAMES[table], len(SCADA_HEADERS), ws.max_row)
        ws.cell(1, 13, "◄ Navigation"); ws.cell(1, 13).hyperlink = "#'Navigation'!A1"
        _format_scada(ws)
    # Equipment sheets are built from catalog templates and retain point-level metadata in hidden state below.
    for typ in state.get("equipment_types", []):
        for table in TABLES:
            regs = [r for r in typ.get("registers", []) if r.get("table") == table]
            if not regs: continue
            sheet_name = f"{typ.get('code')}-{table.replace('_', ' ').title()}"[:31]
            ws = wb.create_sheet(sheet_name)
            for c, h in enumerate(EQUIPMENT_HEADERS, 1): ws.cell(1, c, h)
            for r, reg in enumerate(regs, 2):
                values = [reg.get("label", ""), "", "", reg.get("plc_table", ""), reg.get("address"), reg.get("access", ""), reg.get("type", ""), reg.get("scale", ""), reg.get("unit", "")]
                for c, v in enumerate(values, 1): ws.cell(r, c, v)
            table_name = re.sub(r"[^A-Za-z0-9_]", "_", f"{typ.get('code')}_{ {'coils':'COIL','discrete_inputs':'DI','input_registers':'IR','holding_registers':'HR'}[table]}")
            _add_table(ws, table_name, len(EQUIPMENT_HEADERS), ws.max_row)
            ws.cell(1, 11, "◄ Navigation"); ws.cell(1, 11).hyperlink = "#'Navigation'!A1"
    meta = wb.create_sheet("_GaiaMetadata")
    meta.sheet_state = "hidden"
    packed = json.dumps(state, ensure_ascii=False, default=str)
    for row, offset in enumerate(range(0, len(packed), 30000), start=1):
        meta.cell(row, 1, packed[offset:offset + 30000])
    output = io.BytesIO(); wb.save(output); return output.getvalue()


def _add_table(ws, name: str, columns: int, last_row: int):
    last_row = max(2, last_row)
    from openpyxl.utils import get_column_letter
    ref = f"A1:{get_column_letter(columns)}{last_row}"
    table = Table(displayName=name[:255], ref=ref)
    from openpyxl.worksheet.filters import AutoFilter
    table.autoFilter = AutoFilter(ref=ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleLight14", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    ws.add_table(table)


def _format_scada(ws):
    for col, width in enumerate([50, 35, 34, 16, 16, 15, 20, 14, 15, 25, 14], 1): ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = width
    for formula, color in [('SEARCH("scada to plc",$G2)', "FCE4D6"), ('SEARCH("plc to scada",$G2)', "E2F0D9"), ('SEARCH("two ways",$G2)', "FFF2CC")]:
        ws.conditional_formatting.add(f"A2:K{ws.max_row}", FormulaRule(formula=[f"=IFERROR({formula},0)"], fill=openpyxl.styles.PatternFill("solid", fgColor=color)))
    for col, options in [(7, "Plc To SCADA,SCADA To Plc,Two Ways"), (8, "FUXA,HMI")]:
        dv = DataValidation(type="list", formula1=f'"{options}"', allow_blank=True); ws.add_data_validation(dv); dv.add(f"{openpyxl.utils.get_column_letter(col)}2:{openpyxl.utils.get_column_letter(col)}{ws.max_row}")

