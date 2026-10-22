from __future__ import annotations

import io

import openpyxl
from fastapi.testclient import TestClient

from backend.main import create_app
from backend.v2_type_creation import _replace_st_identifiers


def test_copy_identifier_renaming_preserves_comments_and_string_literals():
    source = "FB_OLD(); (* FB_OLD stays in this comment *)\n// FB_OLD stays here too\nname := 'FB_OLD';"
    assert _replace_st_identifiers(source, {"FB_OLD": "FB_NEW"}) == "FB_NEW(); (* FB_OLD stays in this comment *)\n// FB_OLD stays here too\nname := 'FB_OLD';"


def _client(tmp_path, monkeypatch):
    monkeypatch.delenv("DATABASE_URL", raising=False)
    monkeypatch.setenv("ADMIN_USERNAME", "type-author")
    monkeypatch.setenv("ADMIN_PASSWORD", "type-author-password")
    app = create_app(data_dir=tmp_path, secret_key="type-create-tests-" * 3, cookie_secure=False)
    client = TestClient(app)
    client.__enter__()
    auth = client.post("/api/login", json={"username": "type-author", "password": "type-author-password"})
    assert auth.status_code == 200
    return client, {"X-CSRF-Token": auth.json()["csrf_token"]}


def test_new_equipment_is_reviewed_then_applied_and_excel_roundtrips(tmp_path, monkeypatch):
    client, headers = _client(tmp_path, monkeypatch)
    try:
        project = client.get("/api/v2/project").json()
        payload = {"revision": project["revision"], "name": "NEW_PUMP", "kind": "Equipment",
                   "description": "New equipment", "communication": "Wired I/O", "folder": "Function Blocs/FB",
                   "duts": [{"role": role, "members": ([{"name": "SPARE", "iec_type": "BOOL"}, {"name": "MAX_SPEED", "iec_type": "REAL", "default_value": 12.5, "unit": "rpm"}] if role == "PA" else [{"name": "SPARE", "iec_type": "BOOL"}, {"name": "SCADA_INPUT_TEST", "iec_type": "BOOL"}] if role == "COM" else [{"name": "SPARE", "iec_type": "BOOL"}, {"name": "ENABLE", "iec_type": "BOOL"}, {"name": "ACTIVE", "iec_type": "BOOL"}] if role == "INT" else [{"name": "SPARE", "iec_type": "BOOL"}])} for role in ["PA", "COM", "EXC", "IO", "INT"]],
                   "processing_code": "IF INTERN.ENABLE THEN\n    INTERN.ACTIVE := TRUE;\nEND_IF", "supervision": [{"member_name": "COM.SCADA_INPUT_TEST"}]}
        no_csrf = client.post("/api/v2/types/preview", json=payload)
        assert no_csrf.status_code == 403
        preview = client.post("/api/v2/types/preview", headers=headers, json=payload)
        assert preview.status_code == 200, preview.text
        result = preview.json()
        assert result["errors"] == [], result["errors"]
        assert "DUT_PA_NEW_PUMP" in result["objects"] and "FB_INST_NEW_PUMP" in result["objects"]
        assert client.get("/api/v2/project").json()["revision"] == project["revision"]
        applied = client.post("/api/v2/types/apply", headers=headers, json={"token": result["token"]})
        assert applied.status_code == 200, applied.text
        assert applied.json()["type"]["name"] == "NEW_PUMP"
        assert client.post("/api/v2/types/apply", headers=headers, json={"token": result["token"]}).status_code == 404
        project = client.get("/api/v2/project").json()
        created = next(t for t in project["types"] if t["name"] == "NEW_PUMP")
        assert {d["role"] for d in created["duts"]} == {"PA", "COM", "EXC", "IO", "INT"}
        assert created["fb"]["body"].startswith("// INPUTS") and "INTERN.ACTIVE := TRUE" in created["fb"]["body"]
        assert created["templates"]["TEMPLATE_PA"][0]["Default value"] == 12.5
        assert created["templates"]["TEMPLATE_TCP"][0]["GVL_COM variable"].endswith("SCADA_INPUT_TEST")
        assert not any(i["type_id"] == "NEW_PUMP" for i in project["instances"])
        workbook_response = client.get("/api/v2/exports/TYPE_NEW_PUMP.xlsx")
        assert workbook_response.status_code == 200
        workbook = openpyxl.load_workbook(io.BytesIO(workbook_response.content), data_only=True)
        assert workbook["TYPE"]["B3"].value == "Equipment"
        headers_row = [c.value for c in workbook["DUT_COM"][1]]
        assert "SPARE" in [workbook["DUT_COM"].cell(r, headers_row.index("Member") + 1).value for r in range(2, workbook["DUT_COM"].max_row + 1)]
    finally:
        client.close()


def test_copy_and_upi_defaults_and_invalid_duplicate_name(tmp_path, monkeypatch):
    client, headers = _client(tmp_path, monkeypatch)
    try:
        project = client.get("/api/v2/project").json()
        preview = client.post("/api/v2/types/preview", headers=headers, json={"revision": project["revision"],
            "name": "UPI_NEW", "kind": "Functional unit (UPI)", "copy_from": "UPI_LOCAL_HMI"})
        assert preview.status_code == 200, preview.text
        data = preview.json()
        assert data["errors"] == [], data["errors"]
        assert data["objects"][-1] == "FB_UPI_NEW"
        new = client.post("/api/v2/types/apply", headers=headers, json={"token": data["token"]})
        assert new.status_code == 200, new.text
        typ = new.json()["type"]
        roles = {d["role"] for d in typ["duts"]}
        assert {"COM", "EXC", "INT"}.issubset(roles)
        assert {"PA", "IO"}.isdisjoint(roles)
        source = next(t for t in project["types"] if t["name"] == "UPI_LOCAL_HMI")
        copied = {d["role"]: d for d in typ["duts"]}
        source_com = next(d for d in source["duts"] if d["role"] == "COM")
        assert len(copied["COM"]["members"]) == len(source_com["members"])
        current = client.get("/api/v2/project").json()
        equipment_copy = client.post("/api/v2/types/preview", headers=headers, json={"revision": current["revision"],
            "name": "DIP1500_COPY", "kind": "Equipment", "copy_from": "DIP1500"}).json()
        assert equipment_copy["errors"] == [], equipment_copy["errors"]
        copied_equipment = client.post("/api/v2/types/apply", headers=headers, json={"token": equipment_copy["token"]}).json()["type"]
        dip = next(t for t in project["types"] if t["name"] == "DIP1500")
        assert {key: len(value) for key, value in copied_equipment["rtu"].items()} == {key: len(value) for key, value in dip["rtu"].items()}
        assert len(copied_equipment["templates"]["TEMPLATE_RTU"]) == len(dip["templates"]["TEMPLATE_RTU"])
        assert "TON_PULSE" in copied_equipment["fb"].get("interface_xml", "")
        latest_revision = client.get("/api/v2/project").json()["revision"]
        duplicate = client.post("/api/v2/types/preview", headers=headers, json={"revision": latest_revision,
            "name": "upi_new", "kind": "Equipment"}).json()
        assert duplicate["token"] is None and any("already exists" in e for e in duplicate["errors"])
        invalid = client.post("/api/v2/types/preview", headers=headers, json={"revision": latest_revision,
            "name": "BAD-NAME", "kind": "Equipment"}).json()
        assert invalid["token"] is None and invalid["errors"]
    finally:
        client.close()
