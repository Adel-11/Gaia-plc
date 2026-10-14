from __future__ import annotations

import re
from pathlib import Path

import openpyxl

from backend.engine import export_excel, export_fuxa, import_excel, import_fuxa, seed_state, validate

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "gaia-io-generator-docs"
XLSM = DOCS / "01_io_list_reference" / "GAIA-ELEC-1451-V1 - IO List.xlsm"
CSV_ORIGINAL = DOCS / "02_fuxa_reference" / "SCADA_V1-devices_1_export_original.csv"
CSV_EXPECTED = DOCS / "02_fuxa_reference" / "SCADA_V1-devices_complet_attendu.csv"


def _normalize_ids(data: bytes) -> bytes:
    text = data.decode("utf-8")
    return re.sub(r"t_[0-9a-f]{8}-[0-9a-f]{8}", "t_NEW-ID-NEW-ID", text).encode("utf-8")


def test_reference_imports_generate_expected_fuxa_csv():
    state = import_excel(XLSM.read_bytes(), {
        "settings": {"project_name": "Gaia Refinery", "document_reference": "GAIA-ELEC-1451", "version": "1", "device_id": "d_abeaaea7-b8914e5b", "device_name": "PLC", "ip": "192.168.0.10", "port": 502, "slave_id": 10, "polling": 500, "type_map": {}},
        "equipment_types": [], "equipment": [], "points": [], "fuxa_template": "", "fuxa_existing": [], "generation_logs": [],
    })
    state = import_fuxa(CSV_ORIGINAL.read_bytes(), state)
    original_ids = {row[2] for row in (line.split(",") for line in CSV_ORIGINAL.read_text(encoding="utf-8").splitlines()) if len(row) > 3 and row[0] == "T@" and row[3].lower() not in ("test holding", "test descrete")}
    generated = export_fuxa(state)
    expected = CSV_EXPECTED.read_bytes()
    # The acceptance case deliberately removes the two manually created TEST tags.
    expected = b"\n".join(line for line in expected.split(b"\n") if not (line.startswith(b"T@,") and b",TEST Holding," in line or line.startswith(b"T@,") and b",Test descrete," in line))
    assert _normalize_ids(generated) == _normalize_ids(expected)
    assert not generated.startswith(b"\xef\xbb\xbf")
    assert b"\r" not in generated
    assert generated.endswith(b"\n\n\n")
    assert all(line.count(b",") == 20 for line in generated.splitlines() if line.startswith(b"T@,"))
    generated_ids = {line.decode("utf-8").split(",")[2] for line in generated.splitlines() if line.startswith(b"T@,")}
    assert original_ids <= generated_ids
    assert b",TEST Holding," not in generated and b",Test descrete," not in generated


def test_excel_roundtrip_keeps_model_and_structured_tables():
    original_state = seed_state()
    generated = export_excel(original_state)
    workbook = openpyxl.load_workbook(__import__("io").BytesIO(generated), data_only=True)
    for ws in workbook.worksheets:
        if ws.title == "Navigation" or ws.sheet_state == "hidden":
            continue
        assert ws.tables
        for table in ws.tables.values():
            ref = table.ref
            assert ref.startswith("A1:")
            headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1) if ws.cell(1, c).value]
            table_headers = [column.name for column in table.tableColumns]
            assert table_headers == headers[:len(table_headers)]
            assert len(table_headers) == len(set(table_headers))
            assert table.autoFilter.ref == ref
            assert int(ref.split(":")[1][1:]) == ws.max_row
    imported = import_excel(generated, {"settings": {}, "equipment_types": [], "equipment": [], "points": [], "fuxa_template": "", "fuxa_existing": [], "generation_logs": []})
    regenerated = export_excel(imported)
    reimported = import_excel(regenerated, {})
    assert len(reimported["points"]) == len(imported["points"])
    fields = ("id", "equipment_id", "plc_variable", "table", "address", "scada_name", "words", "fuxa_id", "target", "direction", "unit")
    assert [{k: p.get(k) for k in fields} for p in reimported["points"]] == [{k: p.get(k) for k in fields} for p in imported["points"]]


def test_excel_visible_edits_apply_and_overlapping_words_are_rejected():
    state = seed_state()
    workbook_bytes = export_excel(state)
    workbook = openpyxl.load_workbook(__import__("io").BytesIO(workbook_bytes))
    ws = workbook["SCADA-Coils"]
    ws["B2"] = "Edited visible point"
    ws["A3"] = None
    ws["B3"] = None
    edited = __import__("io").BytesIO()
    workbook.save(edited)
    updated = import_excel(edited.getvalue(), {})
    assert any(point["scada_name"] == "Edited visible point" for point in updated["points"])
    assert not any(point["table"] == "coils" and point["address"] == 2 for point in updated["points"])

    overlapping = {"points": [
        {"id": "a", "table": "holding_registers", "address": 1, "words": 2, "scada_name": "A", "plc_variable": "A", "fuxa_type": "Float32MLE", "target": "HMI"},
        {"id": "b", "table": "holding_registers", "address": 2, "words": 1, "scada_name": "B", "plc_variable": "B", "fuxa_type": "UInt16", "target": "HMI"},
    ]}
    assert any("collision" in error["message"].lower() for error in validate(overlapping)["errors"])


def test_populated_reimport_keeps_references_and_ids_when_addresses_move():
    state = seed_state()
    type_ids = {item["code"]: item["id"] for item in state["equipment_types"]}
    equipment_ids = {item["tag"]: item["id"] for item in state["equipment"]}
    reimported = import_excel(XLSM.read_bytes(), state)
    assert {item["code"]: item["id"] for item in reimported["equipment_types"]} == type_ids
    assert {item["tag"]: item["id"] for item in reimported["equipment"]} == equipment_ids
    equipment_id_set = {item["id"] for item in reimported["equipment"]}
    assert all(point["equipment_id"] in equipment_id_set for point in reimported["points"])

    enriched = seed_state()
    enriched["points"][0]["comment"] = "Keep programmer note"
    enriched["equipment_types"][0]["datasheet"] = "custom://reference"
    enriched_roundtrip = import_excel(export_excel(enriched), enriched)
    assert enriched_roundtrip["points"][0]["comment"] == "Keep programmer note"
    assert next(item for item in enriched_roundtrip["equipment_types"] if item["id"] == enriched["equipment_types"][0]["id"])["datasheet"] == "custom://reference"

    workbook = openpyxl.load_workbook(__import__("io").BytesIO(export_excel(state)))
    ws = workbook["SCADA-Coils"]
    original = next(point for point in state["points"] if point["table"] == "coils")
    original_fuxa_id = original["fuxa_id"]
    row_num = next(row for row in range(2, ws.max_row + 1) if ws.cell(row, 1).value == original["plc_variable"] and ws.cell(row, 4).value == original["address"])
    ws.cell(row_num, 4).value = 504
    edited = __import__("io").BytesIO(); workbook.save(edited)
    moved = import_excel(edited.getvalue(), {})
    moved_point = next(point for point in moved["points"] if point["plc_variable"] == original["plc_variable"] and point["table"] == "coils")
    assert moved_point["address"] == 504
    assert moved_point["fuxa_id"] == original_fuxa_id


def test_fuxa_device_settings_import_and_excel_navigation_catalog_edits():
    state = seed_state()
    imported_csv = CSV_ORIGINAL.read_text(encoding="utf-8").replace("d_abeaaea7-b8914e5b", "d_new-device-00000000")
    parsed = import_fuxa(imported_csv.encode(), state)
    assert parsed["settings"]["device_id"] == "d_new-device-00000000"
    parsed["settings"]["device_id"] = "d_changed-device-00000000"
    generated = export_fuxa(parsed).decode("utf-8")
    assert "D@,d_changed-device-00000000," in generated
    assert ",PLC Connection Status,PLC Connection Status,number,d_changed-device-00000000," in generated

    workbook = openpyxl.load_workbook(__import__("io").BytesIO(export_excel(state)))
    nav = workbook["Navigation"]
    existing_sheets = set(workbook.sheetnames)
    for row in nav.iter_rows(min_row=3):
        cell = row[1]
        if cell.value in existing_sheets:
            assert cell.hyperlink and str(cell.hyperlink.target).startswith("#")
        else:
            assert cell.hyperlink is None
    assert "DIP1500-Discrete Inputs" not in workbook.sheetnames
    reg_ws = workbook["ADM100W-Input Registers"]
    reg_ws["A2"] = "Edited meter register"
    edited = __import__("io").BytesIO(); workbook.save(edited)
    updated = import_excel(edited.getvalue(), {})
    adm = next(item for item in updated["equipment_types"] if item["code"] == "ADM100W")
    assert adm["registers"][0]["label"] == "Edited meter register"

