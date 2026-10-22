from __future__ import annotations

import copy
import io
from pathlib import Path

import openpyxl

from backend.v2_engine import (
    DEFAULT_PACK,
    export_fuxa,
    export_workbook,
    import_workbook,
    load_pack,
    validate_project,
    workbook_to_data,
)


def _source_bytes(filename: str) -> bytes:
    return (DEFAULT_PACK / "02_excel" / filename).read_bytes()


def _loaded_state():
    return copy.deepcopy(load_pack())


def test_imported_new_type_is_part_of_the_normalized_program_model():
    state = _loaded_state()
    wb = openpyxl.load_workbook(io.BytesIO(_source_bytes('TYPE_DOSENSOR.xlsx')))
    ws = wb['TYPE']
    for row in ws.iter_rows(min_row=2):
        if isinstance(row[1].value, str):
            row[1].value = row[1].value.replace('DOSENSOR', 'NEW_SENSOR')
    stream = io.BytesIO()
    wb.save(stream)
    result = import_workbook(stream.getvalue(), 'TYPE_NEW_SENSOR.xlsx', state)
    assert result['validation']['errors'] == []
    assert any(t['name'] == 'NEW_SENSOR' for t in result['state']['types'])


def test_reference_model_counts_and_uncertain_values_are_nonblocking():
    state = _loaded_state()
    assert len(state["types"]) == 6
    assert len(state["instances"]) == 6
    assert sum(len(t["duts"]) for t in state["types"]) == 24
    assert len(state["exchange_functions"]) == 9
    assert len(state["fuxa"]["points"]) == 99
    assert state["snapshot"]["objects"]
    assert state["types"][0]["fb"]["body"]

    validation = validate_project(state)
    assert validation["errors"] == []
    assert any("TO CHECK" in warning for warning in validation["warnings"])


def test_all_eleven_workbooks_roundtrip_cell_values_headers_and_order():
    state = _loaded_state()
    assert len(state["workbooks"]) == 11
    for filename, original in state["workbooks"].items():
        exported = export_workbook(state, filename)
        actual = workbook_to_data(exported, filename)
        assert [s["name"] for s in actual["sheets"]] == [s["name"] for s in original["sheets"]]
        for before, after in zip(original["sheets"], actual["sheets"]):
            assert before["headers"] == after["headers"], (filename, before["name"])
            assert before["rows"] == after["rows"], (filename, before["name"])


def test_unchanged_merge_roundtrip_preserves_all_eleven_workbooks():
    state = _loaded_state()
    for filename, original in state["workbooks"].items():
        result = import_workbook(_source_bytes(filename), filename, state, mode="merge")
        assert result["validation"]["errors"] == [], (filename, result["validation"]["errors"])
        state = result["state"]
        exported = workbook_to_data(export_workbook(state, filename), filename)
        assert [s["name"] for s in exported["sheets"]] == [s["name"] for s in original["sheets"]]
        for before, after in zip(original["sheets"], exported["sheets"]):
            assert before["headers"] == after["headers"], (filename, before["name"])
            assert before["rows"] == after["rows"], (filename, before["name"])


def test_export_refreshes_lists_for_new_types_instances_and_exchange_functions():
    state = _loaded_state()
    state["types"].append({"id": "NEW_TYPE", "name": "NEW_TYPE"})
    state["instances"].append({"id": "NEW_TYPE_1", "name": "NEW_TYPE_1", "type_id": "NEW_TYPE"})
    state["exchange_functions"].append({"id": "FC_EXC_NEW", "name": "FC_EXC_NEW"})

    output = workbook_to_data(export_workbook(state, "MODBUS_TCP.xlsx"), "MODBUS_TCP.xlsx")
    lists = next(sheet for sheet in output["sheets"] if sheet["name"] == "LISTS")
    for header, expected in (("OBJ_TYPE", "NEW_TYPE"), ("INSTANCE", "NEW_TYPE_1"), ("FC_EXC", "FC_EXC_NEW")):
        column = lists["headers"].index(header)
        assert expected in [row[column] for row in lists["rows"] if column < len(row)]


def test_fuxa_reference_is_exact_with_fast_path_and_forced_generation():
    state = _loaded_state()
    expected = (DEFAULT_PACK / "04_fuxa" / "SCADA_V1-devices_v4.csv").read_bytes()
    assert export_fuxa(state) == expected

    # Clearing the unchanged-output digest exercises CSV regeneration, retaining
    # fixed sections, existing IDs, order, and the two reference test tags.
    state["fuxa"]["tcp_digest"] = "force-regeneration"
    assert export_fuxa(state) == expected


def test_bad_upload_preview_does_not_mutate_current_state():
    state = _loaded_state()
    before = copy.deepcopy(state)
    result = import_workbook(b"not an xlsx", "MODBUS_TCP.xlsx", state)
    assert result["validation"]["errors"]
    assert result["state"] == before
    assert state == before

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "COILS"
    sheet.append([None, "Client"])
    sheet.append([1, "FUXA"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    result = import_workbook(buffer.getvalue(), "MODBUS_TCP.xlsx", state)
    assert any("header" in error.lower() for error in result["validation"]["errors"])
    assert result["state"] == before
    assert state == before


def test_import_detects_tcp_collision_without_applying_preview():
    state = _loaded_state()
    wb = openpyxl.load_workbook(io.BytesIO(_source_bytes("MODBUS_TCP.xlsx")))
    ws = wb["COILS"]
    ws["A3"] = ws["A2"].value
    buffer = io.BytesIO()
    wb.save(buffer)

    before = copy.deepcopy(state)
    result = import_workbook(buffer.getvalue(), "MODBUS_TCP.xlsx", state)
    assert any("duplicate tcp modbus address" in error.lower() for error in result["validation"]["errors"])
    assert result["state"] != before
    assert state == before


def test_tcp_collision_is_detected_across_clients_at_same_modbus_address():
    state = _loaded_state()
    wb = openpyxl.load_workbook(io.BytesIO(_source_bytes("MODBUS_TCP.xlsx")))
    ws = wb["COILS"]
    # HMI's client address is offset to 1001, but address 1 is already occupied
    # in the PLC's shared Modbus table by FUXA.
    ws["A17"] = 1
    ws["B17"] = 1001
    buffer = io.BytesIO()
    wb.save(buffer)
    result = import_workbook(buffer.getvalue(), "MODBUS_TCP.xlsx", state, mode="replace")
    assert any("duplicate tcp modbus address" in error.lower() for error in result["validation"]["errors"])


def test_import_rejects_tcp_point_with_unknown_instance_reference():
    state = _loaded_state()
    wb = openpyxl.load_workbook(io.BytesIO(_source_bytes("MODBUS_TCP.xlsx")))
    ws = wb["COILS"]
    headers = [cell.value for cell in ws[1]]
    instance_column = headers.index("Instance") + 1
    ws.cell(2, instance_column).value = "MISSING_INSTANCE"
    buffer = io.BytesIO()
    wb.save(buffer)

    result = import_workbook(buffer.getvalue(), "MODBUS_TCP.xlsx", state, mode="replace")
    assert any("unknown instance MISSING_INSTANCE" in error for error in result["validation"]["errors"])


def test_import_uses_header_names_and_preserves_unknown_columns():
    state = _loaded_state()
    original_body = next(t for t in state["types"] if t["id"] == "ADM100W")["fb"]["body"]
    wb = openpyxl.load_workbook(io.BytesIO(_source_bytes("TYPE_ADM100W.xlsx")))
    ws = wb["DUT_COM"]
    old_rows = list(ws.iter_rows(values_only=True))
    # Reverse the known columns and add a custom trailing field. Member parsing
    # and export must follow headers, preserving the added data without loss.
    reordered = list(reversed(old_rows[0])) + ["Project extension"]
    ws.delete_rows(1, ws.max_row)
    ws.append(reordered)
    for old in old_rows[1:]:
        ws.append(list(reversed(old)) + ["custom" if old[1] else None])
    buffer = io.BytesIO()
    wb.save(buffer)

    result = import_workbook(buffer.getvalue(), "TYPE_ADM100W.xlsx", state, mode="replace")
    assert result["validation"]["errors"] == []
    imported = result["state"]["workbooks"]["TYPE_ADM100W.xlsx"]
    dut = next(s for s in imported["sheets"] if s["name"] == "DUT_COM")
    assert dut["headers"][0] == "Exchange"
    assert "Project extension" in dut["headers"]

    exported = export_workbook(result["state"], "TYPE_ADM100W.xlsx")
    roundtrip = workbook_to_data(exported, "TYPE_ADM100W.xlsx")
    exported_dut = next(s for s in roundtrip["sheets"] if s["name"] == "DUT_COM")
    assert exported_dut["headers"] == dut["headers"]
    assert exported_dut["rows"] == dut["rows"]
    assert result["state"]["types"][0]["duts"][0]["members"][0]["name"]
    imported_type = next(t for t in result["state"]["types"] if t["id"] == "ADM100W")
    assert imported_type["fb"]["body"] == original_body
    assert imported_type["duts"][0]["name"] == "DUT_COM_ADM100W"


def test_partial_reordered_import_merges_old_rows_and_fields_but_replace_drops_them():
    initial = _loaded_state()
    filename = "TYPE_ADM100W.xlsx"
    current = copy.deepcopy(initial)
    current_dut = next(s for s in current["workbooks"][filename]["sheets"] if s["name"] == "DUT_COM")
    current_dut["headers"].append("Legacy field")
    for row in current_dut["rows"]:
        row.append("preserve me")
    original_member = current_dut["rows"][-1][current_dut["headers"].index("Member")]

    wb = openpyxl.load_workbook(io.BytesIO(_source_bytes(filename)))
    ws = wb["DUT_COM"]
    rows = list(ws.iter_rows(values_only=True))
    headers = list(reversed(rows[0])) + ["Incoming field"]
    ws.delete_rows(1, ws.max_row)
    ws.append(headers)
    # Omit the last member row to test that merge keeps old-only natural keys.
    for old in rows[1:-1]:
        ws.append(list(reversed(old)) + ["new upload"])
    incoming = io.BytesIO()
    wb.save(incoming)

    merged = import_workbook(incoming.getvalue(), filename, current, mode="merge")
    assert merged["validation"]["errors"] == []
    merged_dut = next(s for s in merged["state"]["workbooks"][filename]["sheets"] if s["name"] == "DUT_COM")
    assert merged_dut["headers"] == headers + ["Legacy field"]
    assert any(row[merged_dut["headers"].index("Member")] == original_member for row in merged_dut["rows"])
    assert merged_dut["rows"][0][merged_dut["headers"].index("Legacy field")] == "preserve me"
    assert len(merged_dut["rows"]) == len(current_dut["rows"])

    replaced = import_workbook(incoming.getvalue(), filename, current, mode="replace")
    replaced_dut = next(s for s in replaced["state"]["workbooks"][filename]["sheets"] if s["name"] == "DUT_COM")
    assert replaced_dut["headers"] == headers
    assert "Legacy field" not in replaced_dut["headers"]
    assert len(replaced_dut["rows"]) == len(current_dut["rows"]) - 1

