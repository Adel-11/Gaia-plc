from __future__ import annotations

import io
from pathlib import Path

import openpyxl
from fastapi.testclient import TestClient

from backend.main import create_app
from backend.v2_fb_flow import flow_view


ROOT = Path(__file__).resolve().parents[1]


def _login(client):
    response = client.post("/api/login", json={"username": "object-editor", "password": "object-editor-password"})
    assert response.status_code == 200
    return {"X-CSRF-Token": response.json()["csrf_token"]}


def test_object_editor_preview_apply_and_type_workbook_roundtrip(tmp_path, monkeypatch):
    monkeypatch.delenv("DATABASE_URL", raising=False)
    monkeypatch.setenv("ADMIN_USERNAME", "object-editor")
    monkeypatch.setenv("ADMIN_PASSWORD", "object-editor-password")
    app = create_app(data_dir=tmp_path, secret_key="object-editor-test-" * 4, cookie_secure=False)
    with TestClient(app) as client:
        headers = _login(client)
        project = client.get("/api/v2/project").json()
        typ = next(t for t in project["types"] if t["name"] == "DIP1500")
        view = client.get("/api/v2/objects/type/DIP1500").json()
        assert view["kind"] == "type" and view["revision"] == project["revision"]
        assert [d["role"] for d in view["duts"]] == [d["role"] for d in typ["duts"]]
        instance = client.get("/api/v2/objects/instance/DIP1500_1").json()
        assert instance["kind"] == "instance" and instance["type_id"] == "DIP1500"
        dut = next(d for d in typ["duts"] if d["role"] == "INT")
        members = dut["members"]
        edited = [dict(m) for m in members]
        edited.append({"name": "INTERNAL_EDITOR_FLAG", "iec_type": "BOOL", "default_value": False,
                       "section": "Editor", "unit": "", "comment": "Added in object editor"})
        fb = dict(typ["fb"])
        original_flow = flow_view(fb)
        fb["body"] = str(fb.get("body", "")) + "\n(* editor body *)"
        fb["custom_code"] = "(* preserve this exactly\r\nincluding CRLF *)"
        payload = {"revision": project["revision"], "duts": [
            ({**d, "members": edited} if d["role"] == "INT" else d) for d in typ["duts"]],
            "body": fb["body"], "custom_code": fb["custom_code"]}
        no_csrf = client.post("/api/v2/objects/type/DIP1500/preview", json=payload)
        assert no_csrf.status_code == 403
        preview = client.post("/api/v2/objects/type/DIP1500/preview", headers=headers, json=payload)
        assert preview.status_code == 200 and preview.json()["errors"] == []
        token = preview.json()["token"]
        changed = client.post("/api/v2/objects/type/DIP1500/preview", headers=headers,
                              json={**payload, "revision": project["revision"] + 2})
        assert changed.status_code == 409
        assert client.get("/api/v2/project").json()["revision"] == project["revision"]
        applied = client.post("/api/v2/objects/apply", headers=headers, json={"token": token})
        assert applied.status_code == 200, applied.text
        assert client.post("/api/v2/objects/apply", headers=headers, json={"token": token}).status_code == 404
        after = client.get("/api/v2/objects/type/DIP1500").json()
        dut_after = next(d for d in after["duts"] if d["role"] == "INT")
        assert [m["name"] for m in dut_after["members"]] == [m["name"] for m in members] + ["INTERNAL_EDITOR_FLAG"]
        assert after["fb"]["custom_code"] == fb["custom_code"]
        assert "(* editor body *)" in after["fb"]["processing_code"]
        assert after["fb"]["body"] != fb["body"]
        mapping_signature = lambda rows: [(row.get("source"), row.get("destination"), row.get("section")) for row in rows]
        assert mapping_signature(after["fb"]["mappings"]) == mapping_signature(original_flow["mappings"])
        exported = client.get("/api/v2/exports/TYPE_DIP1500.xlsx")
        assert exported.status_code == 200
        workbook = openpyxl.load_workbook(io.BytesIO(exported.content), data_only=True)
        sheet = workbook["DUT_INT"]
        headers = [cell.value for cell in sheet[1]]
        member_col = headers.index("Member") + 1
        assert "INTERNAL_EDITOR_FLAG" in [sheet.cell(row, member_col).value for row in range(2, sheet.max_row + 1)]


def test_object_editor_rejects_fixed_dut_and_referenced_member_delete(tmp_path, monkeypatch):
    monkeypatch.delenv("DATABASE_URL", raising=False)
    monkeypatch.setenv("ADMIN_USERNAME", "object-editor")
    monkeypatch.setenv("ADMIN_PASSWORD", "object-editor-password")
    app = create_app(data_dir=tmp_path, secret_key="object-editor-test-" * 4, cookie_secure=False)
    with TestClient(app) as client:
        headers = _login(client)
        typ = next(t for t in client.get("/api/v2/project").json()["types"] if t["name"] == "DIP1500")
        comm = next(d for d in typ["duts"] if d["role"] == "COM")
        edited = [dict(d, members=list(d["members"])) for d in typ["duts"]]
        target = next(d for d in edited if d["role"] == "COM")
        target["members"] = target["members"][1:]
        preview = client.post("/api/v2/objects/type/DIP1500/preview", headers=headers,
                              json={"revision": typ.get("revision", 1), "duts": edited})
        assert preview.status_code == 200
        errors = preview.json()["errors"]
        assert any("Cannot delete referenced member" in e for e in errors)
        renamed = [dict(d, name="DUT_RENAMED") if d["role"] == "COM" else d for d in edited]
        invalid = client.post("/api/v2/objects/type/DIP1500/preview", headers=headers,
                              json={"revision": client.get("/api/v2/project").json()["revision"], "duts": renamed}).json()
        assert any("DUT roles and order are fixed" in e or "name is fixed" in e for e in invalid["errors"])
