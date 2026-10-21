"""V8 : green headers with tooltips, drop-down lists everywhere (except Unit), instance templates inside the TYPE files,
real rack (P2-11B from the CODESYS device tree)."""
import sys, re, os, shutil
sys.path.insert(0, '/home/claude/v5')
import openpyxl
from openpyxl.styles import PatternFill, Font
from openpyxl.comments import Comment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
import xlgen as X
import xlgen2 as Y
from xlgen import table_sheet, kv_sheet, readme, DUTS, OBJECTS, TODO, TCP_POINTS, mb_ref

OUT = '/home/claude/v5/excel3'
shutil.rmtree(OUT, ignore_errors=True); os.makedirs(f'{OUT}/templates')
GREEN = PatternFill('solid', fgColor='2E7D32')
INST_PH = '<INSTANCE>'

# ======================================================================== real rack (CODESYS device tree, P2_11B)
RACK = [(1, 'P2-16AD-1', 'AI', 16, 'Module', '16 analog inputs, current'),
        (2, 'P2-16DA-1', 'AO', 16, 'Module_3', '16 analog outputs, current'),
        (3, 'P2-32TD2P', 'DO', 32, 'Module_1', '32 discrete outputs 24 VDC, sourcing, 0.2 A per point'),
        (4, 'P2-32ND3-1', 'DI', 32, 'Module_2', '32 discrete inputs 12-24 VDC, sinking / sourcing')] + \
       [(s, None, 'EMPTY', 0, None, 'Empty slot') for s in range(5, 12)]
IO_USED = {(4, 1): 'IOMAP_DI_FAULT_DOSENSOR_1', (3, 1): 'IOMAP_DO_HIGH_DOSENSOR_1', (3, 2): 'IOMAP_DO_LOW_DOSENSOR_1',
           (1, 1): 'IOMAP_AI_SIGNAL_DOSENSOR_1'}
IO_HEAD = ['Slot', 'Module', 'CODESYS device', 'I/O type', 'Channel', 'Tag', 'Description', 'Instance', 'GVL_IOMAP variable', 'IOMAP type',
           'GVL_IO member', 'Signal', 'Range min', 'Range max', 'Unit', 'ZIPLink', 'Status', 'Comment']
RACK_HEAD = ['Slot', 'Module', 'CODESYS device', 'Type', 'Channels', 'Description', 'Status']

def io_row(s, m, dev, t, ch, var, inst_override=None):
    obj, mem, desc = X.iomap_info(var) if var else (None, None, None)
    ioty = X.IOMAP_TYPE.get(var, 'BOOL') if var else None
    return (s, m, dev, t, ch, None, desc, obj, var and f'GVL_IOMAP.{var}', ioty, mem, X.SIGNAL[t], None, None, None, X.ZIPLINK[t],
            TODO if var else 'Spare', 'Channel assignment to confirm' if var else None)

def io_workbook(template=False):
    wb = openpyxl.Workbook()
    readme(wb, 'Physical I/O', [
        'Wired I/O of the PLC rack, one row per channel of each module.', '',
        'Sheets:',
        '  RACK : P2-11B base slot by slot, with the module name in the CODESYS device tree',
        '  IO   : one row per channel. An empty "GVL_IOMAP variable" = spare channel', '',
        'Rules:',
        '  GVL_IOMAP variable : variable mapped on the channel in the CODESYS I/O mapping of the module',
        '  GVL_IO member      : variable of the instance, copied by GFC_IO_IN / GFC_IO_OUT',
        '  New instance : copy the TEMPLATE_IO rows of its TYPE file, replace <INSTANCE>, fill Slot / Channel',
        '  Rack from the CODESYS device tree (P2_11B). Channels used by DOSENSOR_1 not confirmed.'])
    table_sheet(wb, 'RACK', RACK_HEAD, [] if template else [(s, m, d, t, n, desc, 'Installed' if m else 'Empty') for s, m, t, n, d, desc in RACK])
    rows = [] if template else [io_row(s, m, d, t, ch, IO_USED.get((s, ch))) for s, m, t, n, d, _ in RACK for ch in range(1, n + 1)]
    table_sheet(wb, 'IO', IO_HEAD, rows)
    return wb

# ======================================================================== instance templates inside the TYPE files
TPL_SHEETS = ['TEMPLATE_PA', 'TEMPLATE_RTU', 'TEMPLATE_EXC', 'TEMPLATE_TCP', 'TEMPLATE_IO']
TCP_HEAD = [('Instance' if h == 'Object' else h) for h in X.TCP_HEAD]
TCP_SHEET_OF = {'COIL': 'COILS', 'DI': 'DISCRETE_INPUTS', 'IR': 'INPUT_REGISTERS', 'HR': 'HOLDING_REGISTERS'}

def generic(text, inst):
    return text.replace(inst, INST_PH) if isinstance(text, str) else text

def add_instance_templates(wb, o, template=False):
    inst, typ = (o['inst'], o['typ']) if not template else (None, None)
    # ---- PA
    rows = []
    if not template and o['kind'] == 'Equipment':
        for m in DUTS[f'DUT_PA_{typ}']:
            sup = 'Yes' if typ == 'DIP1500' and m['name'] != 'SPARE' else 'No'
            rows.append((m['name'], m['type'], m['init'], None, X.PA_UNIT.get(m['name']), sup, m['comment'] or None, f'GVL_PA.PA_{INST_PH}.{m["name"]}'))
    table_sheet(wb, 'TEMPLATE_PA', Y.PA_HEAD, rows)
    # ---- RTU
    rows = []
    if not template:
        for sh in ('RTU_COILS', 'RTU_DI', 'RTU_IR', 'RTU_HR'):
            for r in X.RTU.get(typ, {}).get(sh, []):
                rows.append((Y.TABLE_OF[sh], r[0], r[1], r[2], r[3], r[4], r[5], r[6], r[8], r[9], r[10], f'GVL_COM.COM_{INST_PH}.{r[11]}', r[12], r[13]))
    table_sheet(wb, 'TEMPLATE_RTU', Y.RTU_I_HEAD, rows)
    # ---- EXC (destination = this type, or source = this type for a UPI)
    rows = []
    if not template:
        for (a, b), lines in Y.exchanges().items():
            for src, sv, dst, dv, g, cmt in lines:
                if dst == inst: rows.append((src, sv, INST_PH, generic(dv, inst), g, cmt))
        if o['kind'] != 'Equipment':
            rows = [(INST_PH, f'GVL_EXC.EXC_{INST_PH}.UPI_OUTPUT_LOCAL_ACTIVE', '<DESTINATION>', '<destination variable>', 'GFC_EXC_IN',
                     'Global local status, one row per equipment instance')]
        for m in DUTS.get(f'DUT_EXC_{typ}', []):          # inputs with no source yet
            if (m['name'].startswith('UPI_INPUT') or m['name'].startswith('EXC_INPUT')) and not any(r[3].endswith('.' + m['name']) for r in rows):
                rows.append(('<SOURCE>', '<source variable>', INST_PH, f'GVL_EXC.EXC_{INST_PH}.{m["name"]}', 'GFC_EXC_IN', m['comment'] or None))
    table_sheet(wb, 'TEMPLATE_EXC', Y.EXC_HEAD, rows)
    # ---- TCP (one row per register, address to give per instance)
    rows = []
    if not template:
        for p in sorted((p for p in TCP_POINTS if p.obj == inst), key=lambda p: (list(TCP_SHEET_OF).index(p.table), p.source, p.addr)):
            client = 'FUXA' if p.source == 'SCADA' else 'HMI'
            d = 'Client to PLC' if p.table in ('COIL', 'HR') else 'PLC to Client'
            for k in range(p.words):
                w = None if p.words == 1 else ('LSW' if k == 0 else 'MSW')
                gc = f'GVL_COM.COM_{INST_PH}.{p.member()}' + (f'.Word_value[{k}]' if p.words == 2 else '')
                rows.append((None, None, client, INST_PH, generic(p.name, inst), d, p.typ, w, X.scale_fix(p.scale), X.unit_fix(p.unit), None, gc,
                             f'Paste in {TCP_SHEET_OF[p.table]}'))
    table_sheet(wb, 'TEMPLATE_TCP', TCP_HEAD, rows)
    # ---- IO
    rows = []
    if not template:
        for m in DUTS.get(f'DUT_IO_{typ}', []):
            if m['name'] == 'SPARE': continue
            var = next((v for v in IO_USED.values() if X.iomap_info(v)[1] == f'GVL_IO.IO_{inst}.{m["name"]}'), None)
            out = m['name'].startswith('OUTPUT')
            t = ('DO' if m['type'] == 'BOOL' else 'AO') if out else ('DI' if m['type'] == 'BOOL' else 'AI')
            rows.append((None, None, None, t, None, None, m['comment'] or None, INST_PH, generic(var and f'GVL_IOMAP.{var}', inst),
                         X.IOMAP_TYPE.get(var, 'BOOL') if var else m['type'], f'GVL_IO.IO_{INST_PH}.{m["name"]}', X.SIGNAL[t], None, None, None,
                         X.ZIPLINK[t], 'Used', None))
    table_sheet(wb, 'TEMPLATE_IO', IO_HEAD, rows)
    return wb

def type_workbook(o, template=False):
    wb = Y.type_workbook(o, template)
    ws = wb['README']
    extra = ['', 'Instance templates (copy into the project files, then replace <INSTANCE> by the instance name):',
             '  TEMPLATE_PA  -> new sheet <instance> of PARAMETERS_PA',
             '  TEMPLATE_RTU -> new sheet <instance> of MODBUS_RTU (+ one row in SLAVES)',
             '  TEMPLATE_EXC -> new sheet <source>-<destination> of EXCHANGES_EXC (+ one row in EXCHANGES)',
             '  TEMPLATE_TCP -> rows of MODBUS_TCP, in the sheet named in "Comment", with free addresses of the client range',
             '  TEMPLATE_IO  -> rows of IO_PHYSICAL / IO, on free channels',
             '  <INSTANCE>, <SOURCE>, <DESTINATION> are placeholders recognized by the generator.']
    r0 = ws.max_row + 1
    for i, l in enumerate(extra):
        ws.cell(r0 + i, 1, l).font = X.BOLD if l.endswith(':') else X.BODY
    return add_instance_templates(wb, o, template)

# ======================================================================== header tooltips
H = {
 # generic
 'Comment': 'Free text. Not used by the generator.',
 'Key': 'Name of the setting. Do not change it, the generator reads the value by this key.',
 'Value': 'Value of the setting. Use the drop-down list when there is one.',
 'Unit': 'Engineering unit of the value (free text).',
 'Scale': 'Conversion raw value -> engineering value, as written in the equipment manual. "1" = no conversion.',
 'Name': 'Readable name of the value. Used as tag name in FUXA and as label in the supervisions.',
 'Description': 'Readable description.',
 'Instance': 'Instance name (e.g. DIP1500_1). Must exist in the program : GVL_COM.COM_<instance>, GVL_EXC.EXC_<instance>...\n<INSTANCE> in a TYPE file template = placeholder.',
 'Status': 'State of the row. TO CHECK = value not confirmed yet.',
 # DUT
 'Section': 'Comment block of the DUT. Written as [Section] in the member comment, starts a new group in the declaration.\nEmpty = same section as the row above.',
 'Member': 'Member name in the DUT (CODESYS identifier : letters, digits, _ , no space).\nPrefixes in DUT_COM : COM_ (Modbus RTU), SCADA_ (FUXA), HMI_ (local HMI), WEBVISU_ (WebVisu).',
 'Default value': 'Initial value in the DUT declaration ( := value ). Empty = type default (FALSE, 0).\nIn DUT_PA : default parameter value, used at first download and after Reset origin.',
 'Exchange': 'Where the member is exchanged : Modbus TCP (FUXA / HMI), Modbus RTU, WebVisu, physical I/O, GFC_EXC, GVL_PA, or nothing (internal).',
 # TYPE / RTU
 'Equipment address': 'Register or coil address as written in the equipment manual (hex 0x4001 or decimal).',
 'Channel offset': 'Position in the CODESYS Modbus channel and in the GVL_MODBUS array (0 = first register of the channel).',
 'Function': 'Modbus function code used by the CODESYS channel.',
 'Word': '32-bit value split in 2 registers : MSW = most significant word, LSW = least significant word. Empty for 16-bit and BOOL.',
 'Word order': 'Order of the 2 words in the equipment : ABCD = MSW at the lower address, CDAB = LSW at the lower address.',
 'DUT_COM member': 'Member of DUT_COM_<type> that receives (read) or provides (write) the value, copied by GFC_COM_IN / GFC_COM_OUT.',
 'CODESYS channel': 'How the CODESYS Modbus master channel runs : cyclic, or on the rising edge of a trigger variable mapped in the slave I/O mapping.',
 'Table': 'Modbus table : Coil, DI (discrete input), IR (input register), HR (holding register).',
 'GVL_MODBUS variable': 'Element of GVL_MODBUS mapped in the CODESYS device I/O mapping (Modbus server or slave channel).\nGVL_MODBUS = exchange variables only.',
 'GVL_COM variable': 'Variable of GVL_COM copied to / from GVL_MODBUS by GFC_COM_IN / GFC_COM_OUT.',
 # RTU lines / slaves
 'Line': 'Name of the serial line (RTU_1, RTU_2...). One line = one COM port, every slave of the line has the same speed and format.',
 'CODESYS COM port device': 'Name of the Modbus COM device in the CODESYS device tree (e.g. Modbus_COM_Port).',
 'CODESYS master device': 'Name of the Modbus master (client) device under the COM port in the CODESYS device tree.',
 'PLC port': 'Serial port of the P2CDS-622 used by the line.',
 'PLC role': 'Role of the PLC on the line.',
 'Baud rate': 'Line speed, same on the PLC port and on every slave.',
 'Parity': 'Parity of the serial frame, same on every slave.',
 'Data bits': 'Data bits of the serial frame.',
 'Stop bits': 'Stop bits of the serial frame.',
 'Response timeout (ms)': 'Time the master waits for a slave answer before an error.',
 'Object type': 'Type of the instance = name of its TYPE file (TYPE_<type>.xlsx).',
 'CODESYS slave device': 'Name of the Modbus slave device in the CODESYS device tree (under the master).',
 'Slave address': 'Modbus address of the equipment (1..247), set on the equipment itself.',
 '32-bit word order': 'Order of the 2 words of 32-bit values for this equipment.',
 # PA
 'Instance value': 'Value for this instance when it differs from the default. Empty = default value.\nDeclared in GVL_PA : PA_<instance> : DUT_PA_<type> := (MEMBER := value);',
 'Changed from supervision': 'Yes = the FB copies a value typed in the active local supervision into GVL_PA (persistent).\nNo = only the declaration value is used.',
 'GVL_PA variable': 'Persistent variable : GVL_PA.PA_<instance>.<member> (VAR_GLOBAL PERSISTENT RETAIN).',
 # EXC
 'Source instance': 'Instance that writes the value (its FB writes the source variable).',
 'Source variable': 'Variable written by the source FB : GVL_EXC.EXC_<source>.UPI_OUTPUT_... or EXC_OUTPUT_...',
 'Destination instance': 'Instance that reads the value.',
 'Destination variable': 'Variable read by the destination FB : GVL_EXC.EXC_<destination>.UPI_INPUT_... or EXC_INPUT_...',
 'Copied by': 'GFC that copies source -> destination.\nGFC_EXC_IN : UPI to equipment, before GFB_INST. GFC_EXC_OUT : equipment to equipment, after GFB_INST.',
 'Sheet': 'Name of the exchange sheet : <source>-<destination>, 31 characters max.',
 'Instance A': 'Source instance of the exchange sheet.',
 'Instance B': 'Destination instance of the exchange sheet.',
 # IO
 'Slot': 'Slot of the module on the P2-11B base (1 = first slot right of the CPU).',
 'Module': 'Module reference installed in the slot.',
 'CODESYS device': 'Name of the module in the CODESYS device tree, under P2_11B (e.g. Module_2). The I/O mapping is done there.',
 'I/O type': 'DI = discrete input, DO = discrete output, AI = analog input, AO = analog output.',
 'Channel': 'Channel number on the module (1..16 analog, 1..32 discrete).',
 'Tag': 'Field tag of the device (P&ID / electrical drawing reference).',
 'GVL_IOMAP variable': 'Variable mapped on the channel in the CODESYS I/O mapping of the module. Empty = spare channel.',
 'IOMAP type': 'Type of the GVL_IOMAP variable, as given by the module mapping.',
 'GVL_IO member': 'Member of the instance in GVL_IO (GVL_IO.IO_<instance>.<member>), copied by GFC_IO_IN / GFC_IO_OUT.',
 'Signal': 'Electrical signal on the channel.',
 'Range min': 'Engineering value at the low end of the signal (e.g. at 4 mA).',
 'Range max': 'Engineering value at the high end of the signal (e.g. at 20 mA).',
 'ZIPLink': 'Pre-wired cable and feedthrough module used for the field wiring.',
 'Channels': 'Number of channels of the module.',
 # TCP
 'Modbus address': 'Address in the PLC Modbus TCP server : coils and registers from 1, discrete inputs from 10001.\nGVL_MODBUS index = address - 1 (discrete inputs : address - 10001).',
 'Client address': 'Address typed in the client. FUXA : discrete inputs 100001.., others = Modbus address.',
 'Client': 'Supervision or device that reads / writes the address. Each client has its own address range (sheet CLIENTS).',
 'Role': 'Role of the PLC for this client.',
 'Address range': 'Addresses reserved for the client in each table.',
 'Word order (32-bit)': 'Word order of 32-bit values for this client.',
 'IP address': 'IP address of the client.',
}
H_CTX = {   # (sheet kind, header) overrides
 ('DUT', 'Type'): 'IEC type of the member : elementary type, DUT, function block or ARRAY[a..b] OF type (type it by hand, the list is a guide).',
 ('DUT', 'Comment'): 'Comment of the member in the DUT declaration ( // ... ).',
 ('RTU', 'Type'): 'Type of the value in the equipment. 32-bit values use 2 rows.',
 ('RTU', 'Direction'): 'Equipment to PLC = read by the PLC. PLC to Equipment = written by the PLC.',
 ('TCP', 'Type'): 'Type of the value seen by the client. REAL uses 2 registers.',
 ('TCP', 'Direction'): 'Client to PLC : coils and holding registers written by the client.\nPLC to Client : discrete inputs and input registers read by the client.',
 ('PA', 'Type'): 'Type of the parameter in DUT_PA_<type>.',
 ('PA', 'Member'): 'Member of DUT_PA_<type>.',
 ('PA', 'Default value'): 'Declaration value in DUT_PA_<type> (TYPE file). Same for every instance.',
 ('PA', 'Comment'): 'Comment of the parameter in DUT_PA_<type>.',
 ('RACK', 'Type'): 'Module type : AI, AO, DI, DO, or EMPTY for a free slot.',
 ('IO', 'Description'): 'What the channel does (comment of the GVL_IO member).',
 ('TCP', 'Instance'): 'Instance that owns the value (GVL_COM.COM_<instance>).',
 ('TCP', 'Comment'): 'Free text. In TEMPLATE_TCP : sheet of MODBUS_TCP where the row goes.',
}
def kind_of(sheet):
    if sheet.startswith('DUT_'): return 'DUT'
    if sheet in ('RTU_COILS', 'RTU_DI', 'RTU_IR', 'RTU_HR', 'TEMPLATE_RTU'): return 'RTU'
    if sheet in ('COILS', 'DISCRETE_INPUTS', 'INPUT_REGISTERS', 'HOLDING_REGISTERS', 'TEMPLATE_TCP', 'CLIENTS'): return 'TCP'
    if sheet in ('RACK',): return 'RACK'
    if sheet in ('IO', 'TEMPLATE_IO'): return 'IO'
    return None

# ======================================================================== drop-down lists
IEC = ['BOOL', 'BYTE', 'WORD', 'DWORD', 'LWORD', 'SINT', 'INT', 'DINT', 'LINT', 'USINT', 'UINT', 'UDINT', 'ULINT', 'REAL', 'LREAL', 'TIME', 'STRING']
LISTS = {
 'IEC_DUT': IEC + ['DUT_Union_Real_Word', 'TON', 'TOF', 'TP', 'R_TRIG', 'F_TRIG', 'ARRAY[0..1] OF WORD', 'ARRAY[0..7] OF BOOL'],
 'VALUE_TYPE': ['BOOL', 'WORD', 'INT', 'UINT', 'DINT', 'UDINT', 'REAL', 'UDINT (time mode) / REAL (volume mode)'],
 'EXCHANGE': ['Modbus TCP FUXA Coil', 'Modbus TCP FUXA DI', 'Modbus TCP FUXA IR', 'Modbus TCP FUXA HR', 'Modbus TCP HMI Coil', 'Modbus TCP HMI DI',
              'Modbus TCP HMI IR', 'Modbus TCP HMI HR', 'WebVisu (direct variable access)', 'Modbus RTU (see RTU tables)',
              'Modbus RTU write trigger (slave I/O mapping)', 'Physical I/O (GVL_IOMAP, GFC_IO)', 'Exchange (GFC_EXC)', 'Persistent parameter (GVL_PA)'],
 'FUNCTION': ['FC01 Read Coils', 'FC02 Read Discrete Inputs', 'FC03 Read Holding Registers', 'FC04 Read Input Registers', 'FC05 Write Single Coil',
              'FC06 Write Single Register', 'FC15 Write Multiple Coils', 'FC16 Write Multiple Registers', 'FC23 Read/Write Multiple Registers'],
 'RTU_DIR': ['Equipment to PLC', 'PLC to Equipment'],
 'TCP_DIR': ['Client to PLC', 'PLC to Client'],
 'WORD': ['MSW', 'LSW'],
 'WORD_ORDER': ['ABCD (MSW first)', 'CDAB (LSW first)', 'BADC (byte swap)', 'DCBA (byte and word swap)'],
 'CHANNEL': ['Read, cyclic', 'Write, cyclic', 'Read, rising edge', 'Write, rising edge of COM_OUTPUT_CMD_TRIGGER',
             'Write, rising edge of COM_OUTPUT_CONFIG_TRIGGER', 'Read, application', 'Write, application', TODO],
 'TABLE': ['Coil', 'DI', 'IR', 'HR'],
 'LINE': ['RTU_1', 'RTU_2', 'RTU_3'],
 'PLC_PORT': ['RS-232/485 4-pin', 'RS-232/485 RJ12', TODO],
 'PLC_ROLE': ['Master', 'Slave'],
 'BAUD': ['1200', '2400', '4800', '9600', '19200', '38400', '57600', '115200', TODO],
 'PARITY': ['None', 'Even', 'Odd', TODO],
 'DATA_BITS': ['7', '8', TODO],
 'STOP_BITS': ['1', '2', TODO],
 'TIMEOUT': ['100', '200', '500', '1000', '2000', TODO],
 'OBJ_TYPE': [o['typ'] for o in OBJECTS],
 'INSTANCE': [o['inst'] for o in OBJECTS] + [INST_PH],
 'YESNO': ['Yes', 'No'],
 'COPIED_BY': ['GFC_EXC_IN', 'GFC_EXC_OUT'],
 'SLOT': [str(i) for i in range(1, 12)],
 'MODULE': ['P2-16AD-1', 'P2-16AD-2', 'P2-16DA-1', 'P2-16DA-2', 'P2-08AD-1', 'P2-08DA-1', 'P2-32ND3-1', 'P2-16ND3', 'P2-32NE3', 'P2-32TD1P',
            'P2-32TD2P', 'P2-16TD2P', 'P2-08TRS', 'P2-FILL'],
 'MOD_TYPE': ['AI', 'AO', 'DI', 'DO', 'EMPTY'],
 'IO_TYPE': ['DI', 'DO', 'AI', 'AO'],
 'IOMAP_TYPE': ['BOOL', 'WORD', 'DWORD', 'INT', 'DINT', 'UINT', 'UDINT', 'REAL'],
 'SIGNAL': ['24 VDC', '4-20 mA', '0-20 mA', '0-10 V', '+/-10 V', 'Dry contact'],
 'ZIPLINK': ['ZL-CBL40-2 / ZL-RTB40-1', 'ZL-P2-CBL24-2 / ZL-RTB20-1'],
 'IO_STATUS': ['Used', 'Spare', 'Reserved', TODO],
 'RACK_STATUS': ['Installed', 'Empty', 'Planned', TODO],
 'CODESYS_DEV': ['Module', 'Module_1', 'Module_2', 'Module_3', 'Module_4', 'Module_5', 'Module_6', 'Module_7', 'Module_8', 'Module_9', 'Module_10'],
 'CLIENT': ['FUXA', 'HMI', 'WEBVISU'],
 'TCP_ROLE': ['PLC = server', 'PLC = client'],
 'CLIENT_ORDER': ['CDAB (Float32MLE)', 'CDAB, confirmed on the PV800', 'ABCD (Float32)', TODO],
 'KIND': ['Equipment', 'Functional unit (UPI)'],
 'COMM': ['Modbus RTU', 'Modbus TCP', 'EtherNet/IP', 'Wired I/O', 'None'],
 'INTERFACE': ['RS-485', 'RS-232'],
 'READ_CH': ['Cyclic', 'Rising edge', 'Application'],
 'SECTION_TCP': ['COILS', 'DISCRETE_INPUTS', 'INPUT_REGISTERS', 'HOLDING_REGISTERS'],
}
# (sheet kind or None, header) -> list name ; ('*', header) for every sheet
COL_LIST = {
 ('DUT', 'Type'): ('IEC_DUT', 'warning'), ('DUT', 'Exchange'): ('EXCHANGE', 'stop'),
 ('RTU', 'Type'): ('VALUE_TYPE', 'warning'), ('RTU', 'Function'): ('FUNCTION', 'stop'), ('RTU', 'Direction'): ('RTU_DIR', 'stop'),
 ('RTU', 'Word order'): ('WORD_ORDER', 'stop'), ('RTU', 'CODESYS channel'): ('CHANNEL', 'warning'), ('RTU', 'Table'): ('TABLE', 'stop'),
 ('TCP', 'Type'): ('VALUE_TYPE', 'warning'), ('TCP', 'Direction'): ('TCP_DIR', 'stop'), ('TCP', 'Client'): ('CLIENT', 'warning'),
 ('TCP', 'Role'): ('TCP_ROLE', 'stop'), ('TCP', 'Word order (32-bit)'): ('CLIENT_ORDER', 'warning'),
 ('RACK', 'Type'): ('MOD_TYPE', 'stop'), ('RACK', 'Status'): ('RACK_STATUS', 'stop'),
 ('IO', 'Status'): ('IO_STATUS', 'stop'),
 ('*', 'Word'): ('WORD', 'stop'), ('*', 'Instance'): ('INSTANCE', 'warning'), ('*', 'Source instance'): ('INSTANCE', 'warning'),
 ('*', 'Destination instance'): ('INSTANCE', 'warning'), ('*', 'Instance A'): ('INSTANCE', 'warning'), ('*', 'Instance B'): ('INSTANCE', 'warning'),
 ('*', 'Line'): ('LINE', 'warning'), ('*', 'PLC port'): ('PLC_PORT', 'stop'), ('*', 'PLC role'): ('PLC_ROLE', 'stop'),
 ('*', 'Baud rate'): ('BAUD', 'stop'), ('*', 'Parity'): ('PARITY', 'stop'), ('*', 'Data bits'): ('DATA_BITS', 'stop'),
 ('*', 'Stop bits'): ('STOP_BITS', 'stop'), ('*', 'Response timeout (ms)'): ('TIMEOUT', 'warning'), ('*', 'Object type'): ('OBJ_TYPE', 'warning'),
 ('*', '32-bit word order'): ('WORD_ORDER', 'stop'), ('*', 'Changed from supervision'): ('YESNO', 'stop'), ('*', 'Copied by'): ('COPIED_BY', 'stop'),
 ('*', 'Slot'): ('SLOT', 'stop'), ('*', 'Module'): ('MODULE', 'warning'), ('*', 'CODESYS device'): ('CODESYS_DEV', 'warning'),
 ('*', 'I/O type'): ('IO_TYPE', 'stop'), ('*', 'IOMAP type'): ('IOMAP_TYPE', 'stop'), ('*', 'Signal'): ('SIGNAL', 'warning'),
 ('*', 'ZIPLink'): ('ZIPLINK', 'warning'),
}
NUMERIC = {'Slave address': (1, 247), 'Channel': (1, 32), 'Channel offset': (0, 65535), 'Modbus address': (1, 19999),
           'Client address': (1, 465535), 'Channels': (0, 32)}
KV_LIST = {'Kind': 'KIND', 'Communication': 'COMM', 'Interface': 'INTERFACE', 'Factory baud rate': 'BAUD', '32-bit word order': 'WORD_ORDER',
           'Read channels': 'READ_CH', 'Object type': 'OBJ_TYPE'}
NOT_TABLE = {'README', 'LISTS'}

def finish(wb):
    """Green headers + tooltips + drop-down lists from a hidden LISTS sheet."""
    ls = wb.create_sheet('LISTS'); ls.sheet_state = 'hidden'
    ref = {}
    for c, (name, values) in enumerate(LISTS.items(), 1):
        ls.cell(1, c, name).font = X.BOLD
        for r, v in enumerate(values, 2): ls.cell(r, c, v)
        L = get_column_letter(c); ref[name] = f'LISTS!${L}$2:${L}${len(values) + 1}'
    for ws in wb:
        if ws.title in NOT_TABLE: continue
        ws.data_validations.dataValidation = []
        title = wb['README']['A1'].value or ''
        k = kind_of(ws.title)
        if title.startswith('Modbus RTU') and ws.title not in ('LINES', 'SLAVES'): k = 'RTU'
        if title.startswith('Parameters') or ws.title == 'TEMPLATE_PA': k = 'PA'
        if title.startswith('Exchanges') or ws.title == 'TEMPLATE_EXC': k = 'EXC'
        headers = [c.value for c in ws[1]]
        for i, h in enumerate(headers, 1):
            cell = ws.cell(1, i)
            cell.fill = GREEN
            text = H_CTX.get((k, h)) or H.get(h)
            if text:
                cm = Comment(text, 'Gaia'); cm.width, cm.height = 320, 40 + 16 * (text.count('\n') + len(text) // 55); cell.comment = cm
            L = get_column_letter(i); rng = f'{L}2:{L}2000'
            spec = COL_LIST.get((k, h)) or COL_LIST.get(('*', h))
            if spec:
                dv = DataValidation(type='list', formula1='=' + ref[spec[0]], allow_blank=True, errorStyle=spec[1],
                                    showErrorMessage=True, error='Value not in the list.' + (' Accepted anyway.' if spec[1] == 'warning' else ''))
                ws.add_data_validation(dv); dv.add(rng)
            elif h in NUMERIC:
                lo, hi = NUMERIC[h]
                dv = DataValidation(type='whole', operator='between', formula1=str(lo), formula2=str(hi), allow_blank=True,
                                    showErrorMessage=True, error=f'Whole number {lo}..{hi}')
                ws.add_data_validation(dv); dv.add(rng)
        if headers[:2] == ['Key', 'Value']:                       # key / value sheets : list per row
            for r in range(2, ws.max_row + 1):
                key = ws.cell(r, 1).value
                if key in KV_LIST:
                    dv = DataValidation(type='list', formula1='=' + ref[KV_LIST[key]], allow_blank=True, errorStyle='warning', showErrorMessage=True)
                    ws.add_data_validation(dv); dv.add(f'B{r}')
    return wb

def save(wb, name):
    finish(wb).save(f'{OUT}/{name}'); return name

def rename_obj(wb):
    for ws in wb:
        for c in ws[1]:
            if c.value == 'Object': c.value = 'Instance'
    return wb

if __name__ == '__main__':
    files = []
    for o in OBJECTS:
        files.append(save(type_workbook(o), f'TYPE_{o["typ"]}.xlsx'))
    files.append(save(io_workbook(), 'IO_PHYSICAL.xlsx'))
    files.append(save(rename_obj(X.tcp_workbook()), 'MODBUS_TCP.xlsx'))
    files.append(save(Y.rtu_workbook(), 'MODBUS_RTU.xlsx'))
    files.append(save(Y.pa_workbook(), 'PARAMETERS_PA.xlsx'))
    files.append(save(Y.exc_workbook(), 'EXCHANGES_EXC.xlsx'))
    for wb, n in ((type_workbook(OBJECTS[0], True), 'TEMPLATE_TYPE.xlsx'), (io_workbook(True), 'TEMPLATE_IO_PHYSICAL.xlsx'),
                  (rename_obj(X.tcp_workbook(True)), 'TEMPLATE_MODBUS_TCP.xlsx'), (Y.rtu_workbook(True), 'TEMPLATE_MODBUS_RTU.xlsx'),
                  (Y.pa_workbook(True), 'TEMPLATE_PARAMETERS_PA.xlsx'), (Y.exc_workbook(True), 'TEMPLATE_EXCHANGES_EXC.xlsx')):
        files.append(save(wb, f'templates/{n}'))
    print('\n'.join(files))
