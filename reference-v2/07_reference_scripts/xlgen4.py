"""V9 : exchanges by FC_EXC_<A>_<B> (type pair) called in GFC_EXC (instance pairs).
EXCHANGES_EXC rebuilt, TEMPLATE_EXC of the TYPE files = GFC_EXC call rows. Everything else as V8."""
import sys, re, os, shutil
sys.path.insert(0, '/home/claude/v5')
import openpyxl
from lxml import etree
import xlgen as X, xlgen2 as Y, xlgen3 as Z
from xlgen import table_sheet, readme, OBJECTS
INST_PH = Z.INST_PH

OUT = '/home/claude/v5/excel4'
Z.OUT = OUT
shutil.rmtree(OUT, ignore_errors=True); os.makedirs(f'{OUT}/templates')

# ------------------------------------------------------------------ read the FC / GFC_EXC model from the generated program
NS = 'http://www.plcopen.org/xml/tc6_0200'; q = lambda t: f'{{{NS}}}{t}'
prg = etree.parse('/home/claude/v6/out/Gaia_V7_FC_EXC.xml').getroot()
BODY = {p.get('name'): ''.join(p.find(f'.//{q("ST")}').itertext()) for p in prg.iter(q('pou'))}
TYPE_OF = lambda i: re.sub(r'_\d+$', '', i)
FCS = []        # (fc, typeA, typeB, [(src, dst)])
for name, b in BODY.items():
    m = re.match(r'FC_EXC_(UPI_LOCAL_\w+?|[A-Z0-9]+)_([A-Z0-9_]+)$', name)
    if not m: continue
    ta, tb = m.groups()
    lines = [(s, d) for d, s in re.findall(r'^(EXC_\w+\.\w+)\s*:=\s*(EXC_\w+\.\w+);', b, re.M)]
    FCS.append((name, ta, tb, lines))
CALLS = re.findall(r'(FC_EXC_\w+)\(\s*(\w+)\s*:= GVL_EXC\.EXC_(\w+),\s*(\w+)\s*:= GVL_EXC\.EXC_(\w+)\);', BODY['GFC_EXC'])
DUTS = X.DUTS
def member_info(typ, mem):
    return next((m for m in DUTS.get(f'DUT_EXC_{typ}', []) if m['name'] == mem), {'type': None, 'comment': None})

# ------------------------------------------------------------------ EXCHANGES_EXC
FC_HEAD = ['Sheet', 'FC', 'Object type A', 'Object type B', 'Parameter A', 'Parameter B', 'Description']
CALL_HEAD = ['FC', 'Instance A', 'Instance B', 'Parameter A value', 'Parameter B value', 'Comment']
LINE_HEAD = ['Direction', 'Source', 'Destination', 'Type', 'Comment']
sheet_of = lambda ta, tb: f'{ta}-{tb}'

def exc_workbook(template=False):
    wb = openpyxl.Workbook()
    readme(wb, 'Exchanges (EXC)', [
        'Exchanges between instances. One function FC_EXC_<A>_<B> per pair of object types, called in GFC_EXC once per pair of instances.', '',
        'Sheets:',
        '  FC_EXC  : list of the exchange functions (one per pair of object types) and their VAR_IN_OUT parameters',
        '  CALLS   : calls in GFC_EXC, one row per pair of instances, with the GVL_EXC variables given to the FC',
        '  <A>-<B> : one sheet per FC, one row per copied member (same for every instance pair of these types)', '',
        'Rules:',
        '  FC name   : FC_EXC_<type A>_<type B>, VAR_IN_OUT EXC_<type A> : DUT_EXC_<type A> ; EXC_<type B> : DUT_EXC_<type B>',
        '  Source / Destination : <parameter>.<member>, e.g. EXC_UPI_LOCAL_FUXA.UPI_OUTPUT_LOCAL_ACTIVE',
        '  GFC_EXC is called in PLC_PRG after GFB_UPI and before GFB_INST',
        '  New instance of an existing type : add its rows in CALLS (TEMPLATE_EXC of its TYPE file)',
        '  New pair of types : add a row in FC_EXC and a new <A>-<B> sheet'])
    table_sheet(wb, 'FC_EXC', FC_HEAD, [] if template else [
        (sheet_of(ta, tb), fc, ta, tb, f'EXC_{ta}', f'EXC_{tb}', f'{ta} -> {tb}') for fc, ta, tb, _ in FCS],
        widths={'Sheet': 32, 'FC': 38, 'Description': 40})
    table_sheet(wb, 'CALLS', CALL_HEAD, [] if template else [
        (fc, ia, ib, f'GVL_EXC.EXC_{ia}', f'GVL_EXC.EXC_{ib}', None) for fc, pa, ia, pb, ib in CALLS],
        widths={'FC': 38, 'Parameter A value': 34, 'Parameter B value': 34})
    for fc, ta, tb, lines in ([] if template else FCS):
        rows = []
        for s, d in lines:
            dt, dm = d.split('.', 1)
            mi = member_info(TYPE_OF(dt[4:]), dm)
            direction = 'A -> B' if s.startswith(f'EXC_{ta}.') else 'B -> A'
            rows.append((direction, s, d, mi['type'], mi['comment'] or None))
        table_sheet(wb, sheet_of(ta, tb), LINE_HEAD, rows, widths={'Source': 50, 'Destination': 50, 'Comment': 50})
    if template:
        table_sheet(wb, '<A>-<B>', LINE_HEAD, [], widths={'Source': 50, 'Destination': 50, 'Comment': 50})
    return wb

# ------------------------------------------------------------------ TEMPLATE_EXC of the TYPE files = CALLS rows
def type_workbook(o, template=False):
    wb = Z.type_workbook(o, template)
    i = wb.sheetnames.index('TEMPLATE_EXC'); del wb['TEMPLATE_EXC']
    rows = []
    if not template:
        for fc, ta, tb, _ in FCS:
            if tb == o['typ']:
                a = next(ia for f, pa, ia, pb, ib in CALLS if f == fc)
                rows.append((fc, a, INST_PH, f'GVL_EXC.EXC_{a}', f'GVL_EXC.EXC_{INST_PH}', None))
            elif ta == o['typ']:
                rows.append((fc, INST_PH, '<INSTANCE_B>', f'GVL_EXC.EXC_{INST_PH}', 'GVL_EXC.EXC_<INSTANCE_B>', 'One row per instance B'))
    ws = table_sheet(wb, 'TEMPLATE_EXC', CALL_HEAD, rows, widths={'FC': 38, 'Parameter A value': 34, 'Parameter B value': 34})
    wb.move_sheet(ws, offset=i - wb.sheetnames.index('TEMPLATE_EXC'))
    rd = wb['README']
    for row in rd.iter_rows():
        c = row[0]
        if c.value and 'TEMPLATE_EXC ->' in str(c.value):
            c.value = '  TEMPLATE_EXC -> rows of EXCHANGES_EXC / CALLS (GFC_EXC calls)'
    return wb

# ------------------------------------------------------------------ lists and tooltips
Z.LISTS['FC_EXC'] = [fc for fc, *_ in FCS]
Z.LISTS['EXC_DIR'] = ['A -> B', 'B -> A']
Z.LISTS['INSTANCE'] = Z.LISTS['INSTANCE'] + ['<INSTANCE_B>']
Z.COL_LIST.update({('*', 'FC'): ('FC_EXC', 'warning'), ('*', 'Object type A'): ('OBJ_TYPE', 'warning'),
                   ('*', 'Object type B'): ('OBJ_TYPE', 'warning'), ('EXC', 'Direction'): ('EXC_DIR', 'stop'),
                   ('EXC', 'Type'): ('IEC_DUT', 'warning')})
Z.H.update({
 'FC': 'Exchange function FC_EXC_<type A>_<type B> (folder Functions / FC_EXC). Same FC for every instance pair of these two types.',
 'Object type A': 'First object type of the FC (first VAR_IN_OUT parameter). For UPI exchanges : the UPI type.',
 'Object type B': 'Second object type of the FC (second VAR_IN_OUT parameter).',
 'Parameter A': 'Name of the first VAR_IN_OUT of the FC : EXC_<type A>, type DUT_EXC_<type A>.',
 'Parameter B': 'Name of the second VAR_IN_OUT of the FC : EXC_<type B>, type DUT_EXC_<type B>.',
 'Parameter A value': 'Variable given to parameter A in the GFC_EXC call : GVL_EXC.EXC_<instance A>.',
 'Parameter B value': 'Variable given to parameter B in the GFC_EXC call : GVL_EXC.EXC_<instance B>.',
 'Source': 'Member read by the FC : <parameter>.<member>, e.g. EXC_UPI_LOCAL_FUXA.UPI_OUTPUT_LOCAL_ACTIVE (an OUTPUT of the source).',
 'Destination': 'Member written by the FC : <parameter>.<member>, e.g. EXC_DIP1500.UPI_INPUT_GLOBAL_LOCAL_FUXA (an INPUT of the destination).',
 'Instance A': 'Instance given to parameter A in the GFC_EXC call.',
 'Instance B': 'Instance given to parameter B in the GFC_EXC call.',
 'Sheet': 'Name of the sheet that lists the copies of the FC : <type A>-<type B>, 31 characters max.',
})
Z.H_CTX.update({('EXC', 'Direction'): 'A -> B : the FC copies a member of A into B. B -> A : the other way.',
                ('EXC', 'Type'): 'Type of the copied member (from DUT_EXC).',
                ('EXC', 'Comment'): 'Comment of the destination member in its DUT_EXC.'})

if __name__ == '__main__':
    files = []
    for o in OBJECTS:
        files.append(Z.save(type_workbook(o), f'TYPE_{o["typ"]}.xlsx'))
    files.append(Z.save(Z.io_workbook(), 'IO_PHYSICAL.xlsx'))
    files.append(Z.save(Z.rename_obj(X.tcp_workbook()), 'MODBUS_TCP.xlsx'))
    files.append(Z.save(Y.rtu_workbook(), 'MODBUS_RTU.xlsx'))
    files.append(Z.save(Y.pa_workbook(), 'PARAMETERS_PA.xlsx'))
    files.append(Z.save(exc_workbook(), 'EXCHANGES_EXC.xlsx'))
    for wb, n in ((type_workbook(OBJECTS[0], True), 'TEMPLATE_TYPE.xlsx'), (Z.io_workbook(True), 'TEMPLATE_IO_PHYSICAL.xlsx'),
                  (Z.rename_obj(X.tcp_workbook(True)), 'TEMPLATE_MODBUS_TCP.xlsx'), (Y.rtu_workbook(True), 'TEMPLATE_MODBUS_RTU.xlsx'),
                  (Y.pa_workbook(True), 'TEMPLATE_PARAMETERS_PA.xlsx'), (exc_workbook(True), 'TEMPLATE_EXCHANGES_EXC.xlsx')):
        files.append(Z.save(wb, f'templates/{n}'))
    print('\n'.join(files))
