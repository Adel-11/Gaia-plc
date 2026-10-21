"""Generate the Excel data files (objects, physical I/O, Modbus TCP) + empty templates.
Source of truth : the PLC program (Gaia V5 export + V6 PA changes), the TCP exchange points (points.py),
the RTU tables of the IO List v2."""
import sys, re, os, copy, collections
sys.path.insert(0, '/home/claude/v4')
from points import ALL as TCP_POINTS
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from lxml import etree

PRG = '/home/claude/v5/out/Gaia_V6_full.xml'
IOLIST = '/home/claude/v4/out/GAIA-ELEC-1451-V1 - IO List.xlsx'
OUT = '/home/claude/v5/excel'
os.makedirs(OUT, exist_ok=True)
NS = 'http://www.plcopen.org/xml/tc6_0200'; q = lambda t: f'{{{NS}}}{t}'
root = etree.parse(PRG).getroot()

# ============================================================== styles
F = 'Arial'
H_FILL = PatternFill('solid', fgColor='1F3A5F'); H_FONT = Font(name=F, bold=True, color='FFFFFF', size=10)
K_FILL = PatternFill('solid', fgColor='E8EEF5'); BODY = Font(name=F, size=10); BOLD = Font(name=F, size=10, bold=True)
TODO_FILL = PatternFill('solid', fgColor='FFF2CC')
THIN = Side(style='thin', color='C8D0DA'); BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TODO = 'TO CHECK'

def table_sheet(wb, name, headers, rows, widths=None, lists=None):
    """Header row + data rows, frozen header, autofilter, drop-down lists on some columns."""
    ws = wb.create_sheet(name)
    ws.append(headers)
    for c in ws[1]:
        c.font, c.fill, c.border = H_FONT, H_FILL, BORDER
        c.alignment = Alignment(vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = 30
    for r in rows:
        ws.append(list(r))
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for c in row:
            c.font, c.border = BODY, BORDER
            if c.value == TODO: c.fill = TODO_FILL
    for i, h in enumerate(headers, 1):
        w = (widths or {}).get(h) or min(60, max(12, len(h) + 2, *(len(str(r[i - 1])) + 2 for r in rows if i <= len(r) and r[i - 1] is not None)) if rows else max(14, len(h) + 2))
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(headers))}{max(2, ws.max_row)}'
    for col, values in (lists or {}).items():
        i = headers.index(col) + 1
        dv = DataValidation(type='list', formula1='"' + ','.join(values) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f'{get_column_letter(i)}2:{get_column_letter(i)}2000')
    return ws

def kv_sheet(wb, name, rows):
    """Key / Value / Comment sheet (one parameter per row)."""
    ws = table_sheet(wb, name, ['Key', 'Value', 'Comment'], rows, widths={'Key': 30, 'Value': 34, 'Comment': 70})
    for row in ws.iter_rows(min_row=2, max_col=1):
        for c in row: c.fill, c.font = K_FILL, BOLD
    return ws

def readme(wb, title, lines):
    ws = wb.active; ws.title = 'README'
    ws['A1'] = title; ws['A1'].font = Font(name=F, size=14, bold=True)
    for i, l in enumerate(lines, 3):
        ws.cell(i, 1, l).font = BOLD if l.endswith(':') else BODY
    ws.cell(len(lines) + 4, 1, 'Cells marked "TO CHECK" (yellow) are values not confirmed yet.').font = Font(name=F, size=10, italic=True)
    ws.column_dimensions['A'].width = 130

def save(wb, name):
    wb.save(f'{OUT}/{name}')
    return name

# ============================================================== PLC program model
def typ(v):
    t = v.find(q('type'))[0]; n = t.tag.split('}')[1]
    if n == 'derived': return t.get('name')
    if n == 'array':
        d = t.find(q('dimension')); b = t.find(q('baseType'))[0]
        bn = b.get('name') if b.tag.endswith('derived') else b.tag.split('}')[1]
        return f'ARRAY[{d.get("lower")}..{d.get("upper")}] OF {bn}'
    return n
def doc(v):
    d = v.find(q('documentation')); return ''.join(d.itertext()).strip() if d is not None else ''
def init(v):
    i = v.find(f'{q("initialValue")}/{q("simpleValue")}'); return i.get('value') if i is not None else None

DUTS = {}
for d in root.iter(q('dataType')):
    rows, sec = [], ''
    for v in d.find(f'{q("baseType")}/{q("struct")}').findall(q('variable')):
        c = doc(v); m = re.match(r'\[([^\]]+)\]\s*(.*)', c, re.S)
        if m: sec, c = m.group(1), m.group(2)
        rows.append(dict(section=sec, name=v.get('name'), type=typ(v), init=init(v), comment=' '.join(c.split())))
    DUTS[d.get('name')] = rows
BODIES = {p.get('name'): ''.join(p.find(f'.//{q("ST")}').itertext()) for p in root.iter(q('pou'))}

# GFC copies : member path -> counterpart, with the GFC name
LINKS = collections.defaultdict(list)
for g, body in BODIES.items():
    if not g.startswith('GFC_'): continue
    for left, right in re.findall(r'^\s*(GVL_[\w.\[\]]+)\s*:=\s*(GVL_[\w.\[\]]+);', body, re.M):
        LINKS[left].append((right, g)); LINKS[right].append((left, g))

# ============================================================== objects
OBJECTS = [
    dict(inst='ADM100W_1', typ='ADM100W', fb='FB_INST_ADM100W', kind='Equipment', desc='Trumeter ADM100W power meter',
         call='PLC_PRG.GFB_INST_1.FB_ADM100W_1', comm='Modbus RTU', folder='01-FB_ADM100W'),
    dict(inst='DOSENSOR_1', typ='DOSENSOR', fb='FB_INST_DOSENSOR', kind='Equipment', desc='Dissolved oxygen sensor (wired analog + digital I/O)',
         call='PLC_PRG.GFB_INST_1.FB_DOSENSOR_1', comm='Wired I/O', folder='02-FB_DO_SENSOR'),
    dict(inst='DIP1500_1', typ='DIP1500', fb='FB_INST_DIP1500', kind='Equipment', desc='Kamoer DIP1500 peristaltic pump',
         call='PLC_PRG.GFB_INST_1.FB_DIP1500_1', comm='Modbus RTU', folder='03-FB_KAMOER_DIP1500'),
    dict(inst='UPI_LOCAL_FUXA', typ='UPI_LOCAL_FUXA', fb='FB_UPI_LOCAL_FUXA', kind='Functional unit (UPI)', desc='Global FUXA local mode',
         call='PLC_PRG.GFB_UPI_1.FB_UPI_LOCAL_FUXA_1', comm='None', folder='04-FB_UPI_LOCAL_FUXA'),
    dict(inst='UPI_LOCAL_HMI', typ='UPI_LOCAL_HMI', fb='FB_UPI_LOCAL_HMI', kind='Functional unit (UPI)', desc='Global HMI local mode',
         call='PLC_PRG.GFB_UPI_1.FB_UPI_LOCAL_HMI_1', comm='None', folder='05-FB_UPI_LOCAL_HMI'),
    dict(inst='UPI_LOCAL_WEBVISU', typ='UPI_LOCAL_WEBVISU', fb='FB_UPI_LOCAL_WEBVISU', kind='Functional unit (UPI)', desc='Global WebVisu local mode',
         call='PLC_PRG.GFB_UPI_1.FB_UPI_LOCAL_WEBVISU_1', comm='None', folder='06-FB_UPI_LOCAL_WEBVISU')]
FAMILIES = [('COM', 'GVL_COM', 'COM_'), ('EXC', 'GVL_EXC', 'EXC_'), ('INT', None, None), ('IO', 'GVL_IO', 'IO_'), ('PA', 'GVL_PA', 'PA_')]
FAM_TXT = {'COM': 'Communication : Modbus RTU (COM_), SCADA FUXA (SCADA_), HMI (HMI_), WebVisu (WEBVISU_)',
           'EXC': 'Exchanges with other equipment (EXC_) and functional units (UPI_)',
           'INT': 'Internal variables of the FB (instance variable INTERN)',
           'IO': 'Wired PLC I/O of the object', 'PA': 'Parameters, persistent (GVL_PA, VAR_GLOBAL PERSISTENT RETAIN)'}
PA_UNIT = {'MAX_MOTOR_SPEED': 'rpm', 'CALIBRATION_COEFFICIENT': 'ml/min or rpm'}
TCP_TABLE = {'COIL': 'Coil', 'DI': 'DI', 'IR': 'IR', 'HR': 'HR'}
TCP_BY_MEMBER = {(p.obj, p.member()): p for p in TCP_POINTS}

def link_of(obj, fam, gvl_path, member):
    """Where the member comes from / goes to."""
    if fam == 'COM':
        p = TCP_BY_MEMBER.get((obj['inst'], member))
        if p:
            client = 'FUXA' if p.source == 'SCADA' else 'HMI'
            a = f'{p.addr}' + (f'-{p.addr + 1}' if p.words == 2 else '')
            return f'Modbus TCP {client} {TCP_TABLE[p.table]} {a}'
        if member.startswith('WEBVISU_'): return 'WebVisu (direct variable access)'
        if member.endswith('_TRIGGER'): return 'Modbus RTU write channel trigger (slave I/O mapping)'
    if fam == 'PA': return 'GVL_PA (persistent), changed from the active local supervision' if obj['typ'] == 'DIP1500' else 'GVL_PA (persistent)'
    if gvl_path:
        l = LINKS.get(f'{gvl_path}.{member}')
        if l: return ' ; '.join(f'{x} ({g})' for x, g in l)
    return None

def unit_of(obj, fam, row):
    if fam == 'PA': return PA_UNIT.get(row['name'])
    p = TCP_BY_MEMBER.get((obj['inst'], row['name']))
    if p and p.unit and 'selon' not in str(p.unit): return p.unit
    return None

DUT_HEAD = ['Section', 'Member', 'Type', 'Initial value', 'Unit', 'Comment', 'Link (source / destination)']
DUT_TYPES = ['BOOL', 'BYTE', 'WORD', 'DWORD', 'INT', 'DINT', 'UINT', 'UDINT', 'REAL', 'LREAL', 'TIME', 'STRING', 'DUT_Union_Real_Word']

# ---- RTU (from the IO List v2 equipment sheets)
iol = openpyxl.load_workbook(IOLIST)
def rtu_rows(sheet):
    return [r for r in iol[sheet].iter_rows(min_row=2, values_only=True) if r[0]]
RTU_HEAD = ['Equipment address', 'Channel offset', 'Function', 'Direction', 'Name', 'Type', 'Word', 'Word order', 'Scale', 'Unit',
            'GVL_MODBUS variable', 'DUT_COM member', 'CODESYS channel', 'Comment']
DIRS = ['Equipment to PLC', 'PLC to Equipment']

def word_of(name):
    return 'MSW' if '(MSW)' in name else ('LSW' if '(LSW)' in name else None)
def clean(name): return name.replace(' (MSW)', '').replace(' (LSW)', '')
def unit_fix(u):
    if not u: return None
    return str(u).replace('selon config ADM100W', 'per ADM100W config').replace(' ou ', ' or ').replace('mode temps', 'time mode').replace('mode volume', 'volume mode')
def scale_fix(s):
    if s in (None, '1'): return s
    return str(s).replace('valeur', 'raw').replace('majeur+mineur+revision', 'major+minor+revision')

RTU = {'ADM100W': {}, 'DIP1500': {}}
RTU['ADM100W']['RTU_IR'] = [
    (r[4], r[4], 'FC04 Read Input Registers', 'Equipment to PLC', clean(r[0]), r[6], word_of(r[0]), 'ABCD (MSW first)', scale_fix(r[7]), unit_fix(r[8]),
     f'GVL_MODBUS.ADM100W_InputRegisters[{r[4]}]', f'COM_INPUT_IR[{r[4]}]', 'Read, cyclic', None) for r in rtu_rows('ADM100W-Input Registers')]
RTU['ADM100W']['RTU_COILS'] = [
    (r[4], r[4], 'FC01 Read Coils', 'Equipment to PLC', r[0], 'BOOL', None, None, None, None,
     f'GVL_MODBUS.ADM100W_Alarms[{r[4] // 8}].{r[4] % 8}', f'COM_INPUT_COILS[{r[4] // 8}].{r[4] % 8}', 'Read, cyclic', None) for r in rtu_rows('ADM100W-Coils')]
RTU['DIP1500']['RTU_COILS'] = [
    (r[4], i, 'FC15 Write Multiple Coils', 'PLC to Equipment', r[0], 'BOOL', None, None, None, None,
     f'GVL_MODBUS.DIP1500_Coils[{i}]', f'COM_OUTPUT_COILS[{i}]', TODO, 'Write channel mode not confirmed (cyclic or trigger)')
    for i, r in enumerate(rtu_rows('DIP1500-Coils'))]
RTU['DIP1500']['RTU_IR'] = [
    (r[4], i, 'FC04 Read Input Registers', 'Equipment to PLC', r[0], r[6], None, None, scale_fix(r[7]), unit_fix(r[8]),
     f'GVL_MODBUS.DIP1500_InputRegisters[{i}]', f'COM_INPUT_IR[{i}]', 'Read, cyclic', None) for i, r in enumerate(rtu_rows('DIP1500-Input Registers'))]
hr = []
for i, r in enumerate(rtu_rows('DIP1500-Holding Registers')):
    a = int(r[4], 16)
    if a <= 0x4006: var, mem, fn, ch, k = 'DIP1500_Holding_Cmd', 'COM_OUTPUT_HR_CMD', 'FC16 Write Multiple Registers', 'Write, rising edge of COM_OUTPUT_CMD_TRIGGER', a - 0x4001
    elif a <= 0x400A: var, mem, fn, ch, k = 'DIP1500_Holding_Counters', 'COM_INPUT_HR_COUNTERS', 'FC03 Read Holding Registers', 'Read, cyclic', a - 0x4007
    else: var, mem, fn, ch, k = 'DIP1500_Holding_Config', 'COM_OUTPUT_HR_CONFIG', 'FC16 Write Multiple Registers', 'Write, rising edge of COM_OUTPUT_CONFIG_TRIGGER', a - 0x400B
    t = r[6] if 'selon' not in str(r[6]) else 'UDINT (time mode) / REAL (volume mode)'
    hr.append((r[4], k, fn, 'Equipment to PLC' if r[5] == 'Equipment To Plc' else 'PLC to Equipment', clean(r[0]), t, word_of(r[0]),
               'CDAB (LSW first)' if word_of(r[0]) else None, scale_fix(r[7]), unit_fix(r[8]), f'GVL_MODBUS.{var}[{k}]', f'{mem}[{k}]', ch, None))
RTU['DIP1500']['RTU_HR'] = hr

RTU_CFG_KEYS = [('PLC role', 'Master'), ('Equipment role', 'Slave'), ('Interface', 'RS-485'), ('CODESYS COM port device', 'Modbus_COM_Port'),
                ('CODESYS master device', 'Modbus_Client_COM_Port'), ('CODESYS slave device', None), ('Slave address', None),
                ('Baud rate', None), ('Parity', None), ('Data bits', None), ('Stop bits', None), ('32-bit word order', None),
                ('Response timeout (ms)', None), ('Read channels', 'Cyclic'), ('Write channels', None)]
RTU_CFG = {
    'ADM100W': {'CODESYS slave device': 'ADM100', 'Slave address': TODO, 'Baud rate': 9600, 'Parity': TODO, 'Data bits': TODO, 'Stop bits': TODO,
                '32-bit word order': 'ABCD (MSW first)', 'Response timeout (ms)': TODO, 'Write channels': 'None (read only)'},
    'DIP1500': {'CODESYS slave device': 'Kamoer_DIP1500', 'Slave address': 11, 'Baud rate': 9600, 'Parity': TODO, 'Data bits': TODO, 'Stop bits': TODO,
                '32-bit word order': 'CDAB (LSW first)', 'Response timeout (ms)': TODO,
                'Write channels': 'Holding registers on rising edge triggers (CMD, CONFIG). Coils : TO CHECK'}}
RTU_CMT = {'Slave address': 'DIP1500 : menu C006 = RS485 mode, address set on the pump', 'Baud rate': 'Same for every slave of the line',
           '32-bit word order': 'ABCD = high word at the lower address, CDAB = low word at the lower address'}

def obj_workbook(obj, template=False):
    wb = openpyxl.Workbook()
    readme(wb, 'Object file' if template else f'Object file : {obj["inst"]}', [
        'One file per object (equipment or functional unit). Every object file has the same sheets and the same columns, for import / export.', '',
        'Sheets:',
        '  OBJECT       : identity of the object (instance, FB, GVL variables, communication)',
        '  DUT_COM      : communication data. Prefixes COM_ (Modbus RTU with the equipment), SCADA_ (FUXA), HMI_ (local HMI), WEBVISU_ (WebVisu)',
        '  DUT_EXC      : exchanges with other equipment (EXC_) and with functional units (UPI_)',
        '  DUT_INT      : internal variables of the FB (INTERN)',
        '  DUT_IO       : wired PLC I/O',
        '  DUT_PA       : parameters, persistent. Initial value = default value, kept when changed from a supervision',
        '  RTU_CONFIG   : Modbus RTU line settings (empty values if the object is not on Modbus RTU)',
        '  RTU_COILS, RTU_DI, RTU_IR, RTU_HR : Modbus RTU exchange tables, one row per coil or register', '',
        'Rules:',
        '  One row per DUT member, in declaration order. "Section" starts a new comment block in the DUT.',
        '  Initial value : empty = type default (FALSE, 0).',
        '  RTU tables : one row per address. A 32-bit value uses 2 rows (Word = MSW / LSW).',
        '  An empty sheet keeps its header row.'])
    o = obj if not template else {k: None for k in OBJECTS[0]}
    kv_sheet(wb, 'OBJECT', [
        ('Instance', o['inst'], 'Object name, used in every GVL variable name'),
        ('Object type', o['typ'], 'Suffix of the DUT names (DUT_COM_<type>...)'),
        ('Kind', o['kind'], 'Equipment or Functional unit (UPI)'),
        ('Description', o['desc'], None),
        ('Function block', o['fb'], None),
        ('Instance path', o['call'], 'Where the FB instance is called'),
        ('Project folder', o['folder'], 'Folder under Function Blocs / FB'),
        ('Communication', o['comm'], 'Modbus RTU, Wired I/O, None...'),
        ('GVL_COM variable', f'GVL_COM.COM_{o["inst"]}' if not template else None, 'Type DUT_COM_<type>'),
        ('GVL_EXC variable', f'GVL_EXC.EXC_{o["inst"]}' if not template else None, 'Type DUT_EXC_<type>'),
        ('GVL_IO variable', f'GVL_IO.IO_{o["inst"]}' if not template and o['kind'] == 'Equipment' else None, 'Type DUT_IO_<type>'),
        ('GVL_PA variable', f'GVL_PA.PA_{o["inst"]}' if not template and o['kind'] == 'Equipment' else None, 'Type DUT_PA_<type>, persistent'),
        ('Internal variable', f'{o["call"]}.INTERN' if not template else None, 'Type DUT_INT_<type>')])
    for fam, gvl, pre in FAMILIES:
        rows = []
        if not template:
            dn = f'DUT_{fam}_{obj["typ"]}'
            path = f'{gvl}.{pre}{obj["inst"]}' if gvl else None
            for m in DUTS.get(dn, []):
                rows.append((m['section'] or None, m['name'], m['type'], m['init'], unit_of(obj, fam, m), m['comment'] or None,
                             link_of(obj, fam, path, m['name'])))
            # section only on the first row of each block
            prev = None; out = []
            for r in rows:
                out.append((r[0] if r[0] != prev else None,) + r[1:]); prev = r[0]
            rows = out
        ws = table_sheet(wb, f'DUT_{fam}', DUT_HEAD, rows, widths={'Section': 30, 'Member': 44, 'Type': 24, 'Initial value': 12, 'Unit': 14,
                                                                     'Comment': 60, 'Link (source / destination)': 60})
        ws.cell(1, 1).comment = None
    cfg = RTU_CFG.get(obj['typ']) if not template else None
    kv_sheet(wb, 'RTU_CONFIG', [(k, (cfg.get(k, v) if cfg else None), RTU_CMT.get(k)) for k, v in RTU_CFG_KEYS])
    for sh in ('RTU_COILS', 'RTU_DI', 'RTU_IR', 'RTU_HR'):
        rows = (RTU.get(obj['typ'], {}).get(sh, []) if not template else [])
        table_sheet(wb, sh, RTU_HEAD, rows, widths={'Name': 46, 'GVL_MODBUS variable': 38, 'CODESYS channel': 44, 'Comment': 40, 'Function': 30},
                    lists={'Direction': DIRS, 'Word': ['MSW', 'LSW']})
    return wb

# ============================================================== physical I/O
RACK = [(0, 'P2CDS-622', 'CPU', 0, 'ProductivityCODESYS CPU'),
        (1, 'P2-32ND3-1', 'DI', 32, '32 discrete inputs 12-24 VDC'), (2, 'P2-32ND3-1', 'DI', 32, '32 discrete inputs 12-24 VDC'),
        (3, 'P2-32TD1P', 'DO', 32, '32 discrete outputs 24 VDC'),
        (4, 'P2-16AD-1', 'AI', 16, '16 analog inputs, current'), (5, 'P2-16AD-1', 'AI', 16, '16 analog inputs, current'),
        (6, 'P2-16AD-1', 'AI', 16, '16 analog inputs, current'), (7, 'P2-16AD-1', 'AI', 16, '16 analog inputs, current'),
        (8, 'P2-16DA-1', 'AO', 16, '16 analog outputs, current'), (9, 'P2-16DA-1', 'AO', 16, '16 analog outputs, current'),
        (10, 'P2-FILL', 'FILL', 0, 'Filler, expansion reserve'), (11, 'P2-FILL', 'FILL', 0, 'Filler, expansion reserve')]
SIGNAL = {'DI': '24 VDC', 'DO': '24 VDC', 'AI': '4-20 mA', 'AO': '4-20 mA'}
ZIPLINK = {'DI': 'ZL-CBL40-2 / ZL-RTB40-1', 'DO': 'ZL-CBL40-2 / ZL-RTB40-1', 'AI': 'ZL-P2-CBL24-2 / ZL-RTB20-1', 'AO': 'ZL-P2-CBL24-2 / ZL-RTB20-1'}
# known points : (slot, channel) -> GVL_IOMAP variable ; channel assignment not confirmed
IO_USED = {(1, 1): 'IOMAP_DI_FAULT_DOSENSOR_1', (3, 1): 'IOMAP_DO_HIGH_DOSENSOR_1', (3, 2): 'IOMAP_DO_LOW_DOSENSOR_1', (4, 1): 'IOMAP_AI_SIGNAL_DOSENSOR_1'}
IO_HEAD = ['Slot', 'Module', 'I/O type', 'Channel', 'Tag', 'Description', 'Object', 'GVL_IOMAP variable', 'IOMAP type', 'GVL_IO member',
           'Signal', 'Range min', 'Range max', 'Unit', 'ZIPLink', 'Status', 'Comment']
IO_TYPES = ['DI', 'DO', 'AI', 'AO']

def iomap_info(var):
    """GVL_IOMAP var -> (object, GVL_IO member, description) from the GFC_IO copies."""
    for other, g in LINKS.get(f'GVL_IOMAP.{var}', []):
        m = re.match(r'GVL_IO\.IO_(\w+?_\d+)\.(\w+)', other)
        if m:
            dut = next(d for d in DUTS[f'DUT_IO_{re.sub(r"_\d+$", "", m.group(1))}'] if d['name'] == m.group(2))
            return m.group(1), f'{other}', dut['comment']
    return None, None, None
IOMAP_TYPE = {'IOMAP_AI_SIGNAL_DOSENSOR_1': 'DWORD'}

def io_workbook(template=False):
    wb = openpyxl.Workbook()
    readme(wb, 'Physical I/O', [
        'Wired I/O of the PLC rack, one row per channel of each module.', '',
        'Sheets:',
        '  RACK : modules of the P2-11B base, slot by slot',
        '  IO   : one row per channel. An empty "GVL_IOMAP variable" = spare channel', '',
        'Rules:',
        '  GVL_IOMAP variable : variable mapped on the channel in the CODESYS I/O mapping',
        '  GVL_IO member      : variable of the object, copied by GFC_IO_IN / GFC_IO_OUT',
        '  Status             : Used, Spare, TO CHECK', '',
        'Rack taken from GAIA-ELEC-1416 (9 I/O modules + 2 fillers). Slot order and channel assignment are not confirmed.'])
    rack_rows = [] if template else [(s, m, t, n, d, TODO if 0 < s < 10 else None) for s, m, t, n, d in RACK]
    table_sheet(wb, 'RACK', ['Slot', 'Module', 'Type', 'Channels', 'Description', 'Status'], rack_rows,
                lists={'Type': ['CPU', 'DI', 'DO', 'AI', 'AO', 'FILL']})
    rows = []
    if not template:
        for s, m, t, n, _ in RACK:
            for ch in range(1, n + 1):
                var = IO_USED.get((s, ch))
                obj, mem, desc = iomap_info(var) if var else (None, None, None)
                ioty = IOMAP_TYPE.get(var, 'BOOL') if var else None
                rows.append((s, m, t, ch, None, desc, obj, var and f'GVL_IOMAP.{var}', ioty, mem, SIGNAL[t],
                             None, None, None, ZIPLINK[t], TODO if var else 'Spare', 'Channel assignment to confirm' if var else None))
    table_sheet(wb, 'IO', IO_HEAD, rows, widths={'Description': 44, 'GVL_IOMAP variable': 40, 'GVL_IO member': 46, 'ZIPLink': 26, 'Comment': 30},
                lists={'I/O type': IO_TYPES, 'Status': ['Used', 'Spare', TODO], 'Signal': ['24 VDC', '4-20 mA', '0-20 mA', '0-10 V']})
    return wb

# ============================================================== Modbus TCP
CLIENTS = [('FUXA', 'SCADA (FUXA on the Raspberry Pi)', 'PLC = server', 'Coils 1-1000, DI 10001-11000, IR 1-1000, HR 1-1000', 'CDAB (Float32MLE)', TODO, None),
           ('HMI', 'Local HMI panel (Rockwell PanelView 800)', 'PLC = server', 'FUXA address + 1000', 'CDAB, confirmed on the PV800',
            TODO, 'Enter Key Handshake disabled for numeric entries')]
TCP_SHEETS = [('COIL', 'COILS', 'Coils'), ('DI', 'DISCRETE_INPUTS', 'Discrete inputs'), ('IR', 'INPUT_REGISTERS', 'Input registers'),
              ('HR', 'HOLDING_REGISTERS', 'Holding registers')]
TCP_HEAD = ['Modbus address', 'Client address', 'Client', 'Object', 'Name', 'Direction', 'Type', 'Word', 'Scale', 'Unit',
            'GVL_MODBUS variable', 'GVL_COM variable', 'Comment']
from points import mb_ref
def client_addr(p):
    if p.source == 'SCADA' and p.table == 'DI': return 100000 + p.addr - 10000
    return p.addr

def tcp_workbook(template=False):
    wb = openpyxl.Workbook()
    readme(wb, 'Modbus TCP exchanges', [
        'All Modbus TCP exchanges where the PLC is the server (supervisions, HMI, other devices later).', '',
        'Sheets:',
        '  CLIENTS : one row per client, with its reserved address range',
        '  COILS, DISCRETE_INPUTS, INPUT_REGISTERS, HOLDING_REGISTERS : one row per address used', '',
        'Rules:',
        '  Modbus address : coils and registers from 1, discrete inputs from 10001',
        '  Client address : address typed in the client (FUXA uses 100001.. for discrete inputs)',
        '  A 32-bit REAL uses 2 rows (Word = LSW then MSW, low word first)',
        '  Direction : Client to PLC (coils, holding registers) or PLC to Client (discrete inputs, input registers)',
        '  A new client takes a new address range : add it in CLIENTS first'])
    table_sheet(wb, 'CLIENTS', ['Client', 'Description', 'Role', 'Address range', 'Word order (32-bit)', 'IP address', 'Comment'],
                [] if template else CLIENTS, widths={'Description': 40, 'Address range': 48, 'Word order (32-bit)': 30, 'Comment': 46})
    for tab, sh, _ in TCP_SHEETS:
        rows = []
        if not template:
            for p in sorted((p for p in TCP_POINTS if p.table == tab), key=lambda p: p.addr):
                client = 'FUXA' if p.source == 'SCADA' else 'HMI'
                d = 'Client to PLC' if tab in ('COIL', 'HR') else 'PLC to Client'
                for k in range(p.words):
                    w = (None if p.words == 1 else ('LSW' if k == 0 else 'MSW'))
                    gc = f'GVL_COM.COM_{p.obj}.{p.member()}' + (f'.Word_value[{k}]' if p.words == 2 else '')
                    rows.append((p.addr + k, client_addr(p) + k, client, p.obj, p.name, d, p.typ, w, scale_fix(p.scale), unit_fix(p.unit),
                                 mb_ref(tab, p.addr, k), gc, None))
        table_sheet(wb, sh, TCP_HEAD, rows, widths={'Name': 44, 'GVL_MODBUS variable': 44, 'GVL_COM variable': 64, 'Object': 18},
                    lists={'Client': ['FUXA', 'HMI'], 'Direction': ['Client to PLC', 'PLC to Client'], 'Word': ['LSW', 'MSW'],
                           'Type': ['BOOL', 'WORD', 'INT', 'UINT', 'REAL', 'DINT', 'UDINT']})
    return wb

if __name__ == '__main__':
    # ============================================================== write
    files = []
    for o in OBJECTS:
        files.append(save(obj_workbook(o), f'OBJ_{o["inst"]}.xlsx'))
    files.append(save(io_workbook(), 'IO_PHYSICAL.xlsx'))
    files.append(save(tcp_workbook(), 'MODBUS_TCP.xlsx'))
    os.makedirs(f'{OUT}/templates', exist_ok=True)
    for wb, n in ((obj_workbook(OBJECTS[0], True), 'TEMPLATE_OBJ.xlsx'), (io_workbook(True), 'TEMPLATE_IO_PHYSICAL.xlsx'),
                  (tcp_workbook(True), 'TEMPLATE_MODBUS_TCP.xlsx')):
        wb.save(f'{OUT}/templates/{n}'); files.append('templates/' + n)
    print('\n'.join(files))
