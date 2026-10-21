"""Excel data files, split type / instance.
TYPE_<type>.xlsx (one per object type) + IO_PHYSICAL, MODBUS_TCP, MODBUS_RTU, PARAMETERS_PA, EXCHANGES_EXC (instances) + templates."""
import sys, re, os, shutil, collections
sys.path.insert(0, '/home/claude/v5')
import openpyxl
import xlgen as X
from xlgen import (table_sheet, kv_sheet, readme, DUTS, LINKS, OBJECTS, FAMILIES, RTU, TCP_BY_MEMBER, TODO, PA_UNIT, TCP_TABLE,
                   DUT_HEAD, unit_of, io_workbook, tcp_workbook)

OUT = '/home/claude/v5/excel2'
shutil.rmtree(OUT, ignore_errors=True); os.makedirs(f'{OUT}/templates')
def save(wb, name):
    wb.save(f'{OUT}/{name}'); return name

TYPES = [o for o in OBJECTS]                      # one instance per type for now
INST = {o['inst']: o for o in OBJECTS}
RTU_INST = ['ADM100W_1', 'DIP1500_1']
EQUIP = [o['inst'] for o in OBJECTS if o['kind'] == 'Equipment']

# ================================================================== TYPE files
TYPE_DUT_HEAD = ['Section', 'Member', 'Type', 'Default value', 'Unit', 'Comment', 'Exchange']
def exchange_of(o, fam, m):
    """Type-level description of where a member is exchanged (no instance address)."""
    n = m['name']
    if fam == 'COM':
        p = TCP_BY_MEMBER.get((o['inst'], n))
        if p: return f'Modbus TCP {"FUXA" if p.source == "SCADA" else "HMI"} {TCP_TABLE[p.table]}'
        if n.startswith('WEBVISU_'): return 'WebVisu (direct variable access)'
        if n.endswith('_TRIGGER'): return 'Modbus RTU write trigger (slave I/O mapping)'
        if n.startswith('COM_'): return 'Modbus RTU (see RTU tables)'
    if fam == 'IO':
        if n == 'SPARE': return None
        return 'Physical I/O (GVL_IOMAP, GFC_IO)'
    if fam == 'EXC':
        if n.startswith('UPI_') or n.startswith('EXC_'): return 'Exchange (GFC_EXC)'
    if fam == 'PA' and n != 'SPARE': return 'Persistent parameter (GVL_PA)'
    return None

TYPE_RTU_HEAD = ['Equipment address', 'Channel offset', 'Function', 'Direction', 'Name', 'Type', 'Word', 'Word order', 'Scale', 'Unit',
                 'DUT_COM member', 'CODESYS channel', 'Comment']
RTU_DEF_KEYS = [('Interface', 'RS-485'), ('Factory slave address', None), ('Factory baud rate', None), ('Supported baud rates', None),
                ('32-bit word order', None), ('Read channels', 'Cyclic'), ('Write channels', None), ('Setup on the equipment', None)]
RTU_DEF = {
    'ADM100W': {'Factory slave address': TODO, 'Factory baud rate': TODO, 'Supported baud rates': TODO, '32-bit word order': 'ABCD (MSW first)',
                'Write channels': 'None (read only)', 'Setup on the equipment': TODO},
    'DIP1500': {'Factory slave address': TODO, 'Factory baud rate': TODO, 'Supported baud rates': 'Index 0..7 : 1200 .. 115200 (3 = 9600)',
                '32-bit word order': 'CDAB (LSW first)', 'Write channels': 'Holding registers on rising edge triggers (CMD, CONFIG). Coils : TO CHECK',
                'Setup on the equipment': 'Menu C006 = RS485 mode, address 1..31'}}

def type_workbook(o, template=False):
    wb = openpyxl.Workbook()
    readme(wb, 'Object type' if template else f'Object type : {o["typ"]}', [
        'One file per object type. Everything here is the same for every instance of the type.',
        'Instance data (name, addresses, slave address, parameter values, exchanges) is in the project files :',
        '  IO_PHYSICAL, MODBUS_TCP, MODBUS_RTU, PARAMETERS_PA, EXCHANGES_EXC', '',
        'Sheets:',
        '  TYPE         : identity of the type (FB, DUT names, kind, communication)',
        '  DUT_COM      : communication data. Prefixes COM_ (Modbus RTU), SCADA_ (FUXA), HMI_ (local HMI), WEBVISU_ (WebVisu)',
        '  DUT_EXC      : exchanges with other instances (EXC_) and functional units (UPI_)',
        '  DUT_INT      : internal variables of the FB (INTERN)',
        '  DUT_IO       : wired PLC I/O',
        '  DUT_PA       : parameters. "Default value" = declaration value of the DUT',
        '  RTU_DEFAULT  : Modbus RTU properties of the equipment (empty if not on Modbus RTU)',
        '  RTU_COILS, RTU_DI, RTU_IR, RTU_HR : Modbus RTU register map of the equipment', '',
        'Rules:',
        '  One row per DUT member, in declaration order. "Section" starts a new comment block in the DUT.',
        '  Default value : empty = type default (FALSE, 0).',
        '  RTU tables : one row per address. A 32-bit value uses 2 rows (Word = MSW / LSW).',
        '  An empty sheet keeps its header row.'])
    e = {k: None for k in o} if template else o
    kv_sheet(wb, 'TYPE', [
        ('Object type', e['typ'], 'Suffix of the DUT names'),
        ('Kind', e['kind'], 'Equipment or Functional unit (UPI)'),
        ('Description', e['desc'].replace(' (wired analog + digital I/O)', '') if e['desc'] else None, None),
        ('Function block', e['fb'], None),
        ('Project folder', e['folder'], 'Folder under Function Blocs / FB'),
        ('Communication', e['comm'], 'Modbus RTU, Wired I/O, None'),
        ('DUT_COM', f'DUT_COM_{e["typ"]}' if e['typ'] else None, None), ('DUT_EXC', f'DUT_EXC_{e["typ"]}' if e['typ'] else None, None),
        ('DUT_INT', f'DUT_INT_{e["typ"]}' if e['typ'] else None, None),
        ('DUT_IO', f'DUT_IO_{e["typ"]}' if e['typ'] and o['kind'] == 'Equipment' else None, None),
        ('DUT_PA', f'DUT_PA_{e["typ"]}' if e['typ'] and o['kind'] == 'Equipment' else None, None)])
    for fam, _, _ in FAMILIES:
        rows = []
        if not template:
            prev = None
            for m in DUTS.get(f'DUT_{fam}_{o["typ"]}', []):
                rows.append((m['section'] if m['section'] != prev else None, m['name'], m['type'], m['init'], unit_of(o, fam, m),
                             m['comment'] or None, exchange_of(o, fam, m)))
                prev = m['section']
        table_sheet(wb, f'DUT_{fam}', TYPE_DUT_HEAD, rows, widths={'Section': 30, 'Member': 44, 'Type': 24, 'Default value': 13, 'Unit': 14,
                                                                     'Comment': 60, 'Exchange': 40})
    d = RTU_DEF.get(o['typ']) if not template else None
    kv_sheet(wb, 'RTU_DEFAULT', [(k, (d.get(k, v) if d else (v if not template else None))) for k, v in RTU_DEF_KEYS] if True else [])
    for sh in ('RTU_COILS', 'RTU_DI', 'RTU_IR', 'RTU_HR'):
        rows = [] if template else [r[:10] + r[11:] for r in RTU.get(o['typ'], {}).get(sh, [])]   # without the GVL_MODBUS variable
        table_sheet(wb, sh, TYPE_RTU_HEAD, rows, widths={'Name': 46, 'Function': 30, 'CODESYS channel': 44, 'Comment': 40},
                    lists={'Direction': X.DIRS, 'Word': ['MSW', 'LSW']})
    return wb

# ================================================================== MODBUS_RTU (one sheet per instance)
LINE_HEAD = ['Line', 'CODESYS COM port device', 'CODESYS master device', 'PLC port', 'PLC role', 'Baud rate', 'Parity', 'Data bits', 'Stop bits',
             'Response timeout (ms)', 'Comment']
LINES = [('RTU_1', 'Modbus_COM_Port', 'Modbus_Client_COM_Port', TODO, 'Master', 9600, TODO, TODO, TODO, TODO, None)]
SLAVE_HEAD = ['Instance', 'Object type', 'Line', 'CODESYS slave device', 'Slave address', '32-bit word order', 'Comment']
SLAVES = [('ADM100W_1', 'ADM100W', 'RTU_1', 'ADM100', TODO, 'ABCD (MSW first)', None),
          ('DIP1500_1', 'DIP1500', 'RTU_1', 'Kamoer_DIP1500', 11, 'CDAB (LSW first)', None)]
RTU_I_HEAD = ['Table', 'Equipment address', 'Channel offset', 'Function', 'Direction', 'Name', 'Type', 'Word', 'Scale', 'Unit',
              'GVL_MODBUS variable', 'GVL_COM variable', 'CODESYS channel', 'Comment']
TABLE_OF = {'RTU_COILS': 'Coil', 'RTU_DI': 'DI', 'RTU_IR': 'IR', 'RTU_HR': 'HR'}

def rtu_workbook(template=False):
    wb = openpyxl.Workbook()
    readme(wb, 'Modbus RTU', [
        'Modbus RTU exchanges of the project, the PLC is master.', '',
        'Sheets:',
        '  LINES   : one row per serial line (COM port settings, shared by every slave of the line)',
        '  SLAVES  : one row per instance on Modbus RTU (line, slave device, address)',
        '  <instance> : one sheet per instance, one row per coil or register, with its GVL_MODBUS and GVL_COM variables', '',
        'Rules:',
        '  The register map comes from the TYPE file. This file adds what belongs to the instance : variables and channel settings.',
        '  A 32-bit value uses 2 rows (Word = MSW / LSW).',
        '  Sheet name = instance name, as in SLAVES.'])
    table_sheet(wb, 'LINES', LINE_HEAD, [] if template else LINES, widths={'CODESYS COM port device': 26, 'CODESYS master device': 26})
    table_sheet(wb, 'SLAVES', SLAVE_HEAD, [] if template else SLAVES, widths={'CODESYS slave device': 24, '32-bit word order': 22})
    for inst in (['<INSTANCE>'] if template else RTU_INST):
        rows = []
        if not template:
            o = INST[inst]
            for sh in ('RTU_COILS', 'RTU_DI', 'RTU_IR', 'RTU_HR'):
                for r in RTU.get(o['typ'], {}).get(sh, []):
                    # r : addr, offset, fn, dir, name, type, word, order, scale, unit, gvl_modbus, dut member, channel, comment
                    rows.append((TABLE_OF[sh], r[0], r[1], r[2], r[3], r[4], r[5], r[6], r[8], r[9], r[10], f'GVL_COM.COM_{inst}.{r[11]}', r[12], r[13]))
        table_sheet(wb, inst, RTU_I_HEAD, rows, widths={'Name': 46, 'Function': 30, 'GVL_MODBUS variable': 40, 'GVL_COM variable': 44,
                                                         'CODESYS channel': 44, 'Comment': 40},
                    lists={'Table': ['Coil', 'DI', 'IR', 'HR'], 'Direction': X.DIRS, 'Word': ['MSW', 'LSW']})
    return wb

# ================================================================== PARAMETERS_PA (one sheet per instance)
PA_HEAD = ['Member', 'Type', 'Default value', 'Instance value', 'Unit', 'Changed from supervision', 'Comment', 'GVL_PA variable']
def pa_workbook(template=False):
    wb = openpyxl.Workbook()
    readme(wb, 'Parameters (PA)', [
        'Parameters of each instance. GVL_PA is a persistent variable list (VAR_GLOBAL PERSISTENT RETAIN).', '',
        'Sheets:',
        '  <instance> : one sheet per equipment instance, one row per DUT_PA member', '',
        'Rules:',
        '  Default value  : declaration value of the DUT_PA (TYPE file), same for every instance',
        '  Instance value : value for this instance. Empty = default value.',
        '                   Different from the default : declared in GVL_PA, e.g. PA_DIP1500_1 : DUT_PA_DIP1500 := (CALIBRATION_COEFFICIENT := 3.75);',
        '  Changed from supervision : Yes = the FB copies a new value typed in the active local supervision into PA (kept after restart and download)',
        '  The value stored in the PLC can differ from this file once changed from a supervision. Reset origin = back to declaration values.'])
    for inst in (['<INSTANCE>'] if template else EQUIP):
        rows = []
        if not template:
            o = INST[inst]
            for m in DUTS[f'DUT_PA_{o["typ"]}']:
                sup = 'Yes' if o['typ'] == 'DIP1500' and m['name'] != 'SPARE' else 'No'
                rows.append((m['name'], m['type'], m['init'], None, PA_UNIT.get(m['name']), sup, m['comment'] or None,
                             f'GVL_PA.PA_{inst}.{m["name"]}'))
        table_sheet(wb, inst, PA_HEAD, rows, widths={'Member': 32, 'Comment': 56, 'GVL_PA variable': 46, 'Changed from supervision': 16},
                    lists={'Changed from supervision': ['Yes', 'No']})
    return wb

# ================================================================== EXCHANGES_EXC (one sheet per pair of instances)
EXC_HEAD = ['Source instance', 'Source variable', 'Destination instance', 'Destination variable', 'Copied by', 'Comment']
INDEX_HEAD = ['Sheet', 'Instance A', 'Instance B', 'Description']
def exchanges():
    pairs = collections.OrderedDict()
    body = X.BODIES['GFC_EXC_IN'] + '\n' + X.BODIES.get('GFC_EXC_OUT', '')
    for g in ('GFC_EXC_IN', 'GFC_EXC_OUT'):
        for left, right in re.findall(r'^\s*(GVL_EXC\.[\w.\[\]]+)\s*:=\s*(GVL_EXC\.[\w.\[\]]+);', X.BODIES.get(g, ''), re.M):
            src = re.match(r'GVL_EXC\.EXC_(\w+?)\.(UPI_|EXC_)', right).group(1)
            dst = re.match(r'GVL_EXC\.EXC_(\w+?)\.(UPI_|EXC_)', left).group(1)
            cmt = next((m['comment'] for m in DUTS[f'DUT_EXC_{INST[dst]["typ"]}'] if m['name'] == left.split('.')[-1]), None)
            pairs.setdefault((src, dst), []).append((src, right, dst, left, g, cmt))
    return pairs

def exc_workbook(template=False):
    wb = openpyxl.Workbook()
    readme(wb, 'Exchanges (EXC)', [
        'Exchanges between instances (equipment and functional units), copied by GFC_EXC_IN / GFC_EXC_OUT.', '',
        'Sheets:',
        '  EXCHANGES : list of the exchange sheets',
        '  <A>-<B>   : one sheet per pair of instances, one row per copied variable', '',
        'Rules:',
        '  Source variable : written by the source instance (UPI_OUTPUT_... / EXC_OUTPUT_...)',
        '  Destination variable : read by the destination instance (UPI_INPUT_... / EXC_INPUT_...)',
        '  Copied by : GFC_EXC_IN (UPI to equipment, before GFB_INST) or GFC_EXC_OUT (equipment to equipment, after GFB_INST)',
        '  Sheet name : <source>-<destination>, 31 characters max'])
    pairs = {} if template else exchanges()
    idx = [] if template else [(f'{a}-{b}', a, b, f'{INST[a]["desc"]} -> {INST[b]["desc"].split(" (")[0]}') for a, b in pairs]
    table_sheet(wb, 'EXCHANGES', INDEX_HEAD, idx, widths={'Sheet': 32, 'Description': 70})
    for (a, b), rows in (pairs.items() if not template else [(('<A>', '<B>'), [])]):
        name = f'{a}-{b}'; assert len(name) <= 31, name
        table_sheet(wb, name, EXC_HEAD, rows, widths={'Source variable': 56, 'Destination variable': 56, 'Comment': 46},
                    lists={'Copied by': ['GFC_EXC_IN', 'GFC_EXC_OUT']})
    return wb

# ================================================================== IO / TCP : rename "Object" -> "Instance"
def rename_col(wb):
    for ws in wb:
        for c in ws[1]:
            if c.value == 'Object': c.value = 'Instance'
    return wb

if __name__ == '__main__':
    files = []
    for o in TYPES:
        files.append(save(type_workbook(o), f'TYPE_{o["typ"]}.xlsx'))
    files.append(save(rename_col(io_workbook()), 'IO_PHYSICAL.xlsx'))
    files.append(save(rename_col(tcp_workbook()), 'MODBUS_TCP.xlsx'))
    files.append(save(rtu_workbook(), 'MODBUS_RTU.xlsx'))
    files.append(save(pa_workbook(), 'PARAMETERS_PA.xlsx'))
    files.append(save(exc_workbook(), 'EXCHANGES_EXC.xlsx'))
    for wb, n in ((type_workbook(OBJECTS[0], True), 'TEMPLATE_TYPE.xlsx'), (rename_col(io_workbook(True)), 'TEMPLATE_IO_PHYSICAL.xlsx'),
                  (rename_col(tcp_workbook(True)), 'TEMPLATE_MODBUS_TCP.xlsx'), (rtu_workbook(True), 'TEMPLATE_MODBUS_RTU.xlsx'),
                  (pa_workbook(True), 'TEMPLATE_PARAMETERS_PA.xlsx'), (exc_workbook(True), 'TEMPLATE_EXCHANGES_EXC.xlsx')):
        files.append(save(wb, f'templates/{n}'))
    print('\n'.join(files))
